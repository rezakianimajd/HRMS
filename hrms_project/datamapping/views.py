"""نگاشت و ایمپورت داده — API."""
from datetime import date, datetime

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from accounting.models import (
    AccountGroup, Account, AuxiliaryAccount,
    Journal, FiscalYear, AccountingDocument, AccountingDocumentLine,
)
from datamapping.models import MappingSource, MappingEntry
from datamapping.serializers import MappingSourceSerializer, MappingEntrySerializer
from datamapping.services import (
    normalize, find_best_match, get_candidates, to_decimal,
)
from accounting.coding import suggest_code


def _company(request):
    return getattr(request, 'tenant', None) or getattr(request, 'company', None)


LEVEL_LABELS = {
    'group': 'گروه حساب',
    'general': 'حساب کل',
    'subsidiary': 'حساب معین',
    'auxiliary': 'حساب تفصیلی',
}


def _resolve_target(level, target_id, company):
    if not target_id:
        return None
    if level == 'group':
        return AccountGroup.objects.filter(company=company, id=target_id).first()
    if level == 'general':
        return Account.objects.filter(company=company, id=target_id, parent__isnull=True).first()
    if level == 'subsidiary':
        return Account.objects.filter(company=company, id=target_id).first()
    if level == 'auxiliary':
        return AuxiliaryAccount.objects.filter(company=company, id=target_id).first()
    return None


def _assign_target_fields(level, obj):
    return {
        'group': {'target_group': obj},
        'general': {'target_account': obj},
        'subsidiary': {'target_account': obj},
        'auxiliary': {'target_auxiliary': obj},
    }.get(level, {})


def _match_headers(headers, alias_lists):
    result = {}
    norm_headers = [normalize(h) for h in headers]
    for key, aliases in alias_lists.items():
        for a in aliases:
            na = normalize(a)
            for i, h in enumerate(norm_headers):
                if h == na and key not in result:
                    result[key] = i
                    break
            if key in result:
                break
    return result


# -----------------------------------------------------------------------------
# منابع (Sources)
# -----------------------------------------------------------------------------
@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def sources(request):
    company = _company(request)
    if request.method == 'GET':
        qs = MappingSource.objects.filter(company=company, is_active=True)
        return Response(MappingSourceSerializer(qs, many=True).data)
    serializer = MappingSourceSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    obj = serializer.save(company=company)
    return Response(MappingSourceSerializer(obj).data, status=201)


# -----------------------------------------------------------------------------
# گزینه‌های تطبیق
# -----------------------------------------------------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def options(request):
    company = _company(request)
    level = request.query_params.get('level')

    def acc_brief(acc):
        return {'id': acc.id, 'code': acc.code, 'name': acc.name}

    payload = {
        'groups': [{'id': g.id, 'code': g.code, 'name': g.name} for g in AccountGroup.objects.filter(company=company, is_active=True)],
        'general': [acc_brief(a) for a in Account.objects.filter(company=company, is_active=True, parent__isnull=True)],
        'subsidiary': [acc_brief(a) for a in Account.objects.filter(company=company, is_active=True, parent__isnull=False)],
        'auxiliary': [{'id': a.id, 'code': a.code, 'name': a.name} for a in AuxiliaryAccount.objects.filter(company=company, is_active=True)],
    }
    if level and level in payload:
        return Response(payload[level])
    return Response(payload)


# -----------------------------------------------------------------------------
# پیش‌نمایش تطبیق کدینگ
# -----------------------------------------------------------------------------
CODING_HEADERS = {
    'source_code': ['کد', 'کد حساب', 'کد کل', 'کد مبدا', 'code'],
    'source_name': ['عنوان', 'نام', 'شرح', 'نام حساب', 'title', 'name'],
    'parent_name': ['کد والد', 'کد پدر', 'کد گروه', 'کد نوع', 'والد', 'پدر', 'گروه', 'نوع حساب', 'نوع', 'زیرمجموعه', 'دسته', 'parent'],
}


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def preview(request):
    company = _company(request)
    level = request.data.get('level') or request.query_params.get('level')
    if level not in LEVEL_LABELS:
        return Response({'error': 'سطح کدینگ نامعتبر است'}, status=400)
    kind = request.data.get('kind') or request.query_params.get('kind') or ''

    file = request.FILES.get('file')
    if not file:
        return Response({'error': 'فایل انتخاب نشده است'}, status=400)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        raw = list(ws.iter_rows(values_only=True))
        if not raw:
            return Response({'error': 'فایل خالی است'}, status=400)
        headers = [str(h).strip() if h is not None else '' for h in raw[0]]
        col_map = _match_headers(headers, CODING_HEADERS)
        if 'source_code' not in col_map and 'source_name' not in col_map:
            return Response({'error': 'ستون‌های «کد» یا «عنوان» پیدا نشدند'}, status=400)

        candidates = get_candidates(level, company)
        # برای تفصیلی: فقط کاندیدهای همان دسته (قابل ردیابی دسته به دسته)
        if level == 'auxiliary' and kind:
            cat_name = dict((s['kind'], s['category']) for s in AUX_SOURCES).get(kind)
            if cat_name:
                candidates = [a for a in candidates if getattr(getattr(a, 'category', None), 'name', '') == cat_name]

        existing = {e.source_code: e for e in MappingEntry.objects.filter(company=company, level=level, kind=kind)}

        rows = []
        for r in raw[1:]:
            if not any(r):
                continue
            code = str(r[col_map['source_code']]).strip() if 'source_code' in col_map and col_map['source_code'] < len(r) and r[col_map['source_code']] is not None else ''
            name = str(r[col_map['source_name']]).strip() if 'source_name' in col_map and col_map['source_name'] < len(r) and r[col_map['source_name']] is not None else ''
            parent_name = str(r[col_map['parent_name']]).strip() if 'parent_name' in col_map and col_map['parent_name'] < len(r) and r[col_map['parent_name']] is not None else ''
            if not code and not name:
                continue
            best, score = find_best_match(name, code, candidates)
            prev = existing.get(code)
            rows.append({
                'source_code': code,
                'source_name': name,
                'parent_name': parent_name,
                'matched_target_id': getattr(best, 'id', None),
                'matched_target_code': getattr(best, 'code', ''),
                'matched_target_name': getattr(best, 'name', ''),
                'match_score': round(score, 2),
                'is_exact': bool(best and score >= 0.999),
                'previously_mapped': bool(prev and prev.status in ('matched', 'new')),
                'previous_target_name': prev.target_name if prev else '',
            })
        return Response({'level': level, 'kind': kind, 'rows': rows, 'total': len(rows)})
    except Exception as e:
        return Response({'error': f'خطا در خواندن فایل: {str(e)[:120]}'}, status=400)


# -----------------------------------------------------------------------------
# اعمال نگاشت
# -----------------------------------------------------------------------------
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def apply(request):
    company = _company(request)
    level = request.data.get('level')
    if level not in LEVEL_LABELS:
        return Response({'error': 'سطح کدینگ نامعتبر است'}, status=400)
    source_id = request.data.get('source_id')
    kind = request.data.get('kind') or ''
    rows = request.data.get('rows') or []

    created = 0
    mapped = 0
    ignored = 0
    errors = []

    for i, row in enumerate(rows):
        action = row.get('action', 'ignore')
        source_code = str(row.get('source_code', '') or '').strip()
        source_name = str(row.get('source_name', '') or '').strip()

        if action == 'ignore':
            ignored += 1
            continue

        if action == 'matched':
            target_id = row.get('target_id') or row.get('matched_target_id')
            obj = _resolve_target(level, target_id, company)
            if not obj:
                errors.append({'row': i + 2, 'error': 'هدف انتخاب‌شده یافت نشد'})
                continue
        elif action == 'new':
            name = str(row.get('new_name', '') or source_name or '').strip()
            parent_name = str(row.get('parent_name', '') or '').strip()
            if not name:
                errors.append({'row': i + 2, 'error': 'عنوان جدید خالی است'})
                continue
            obj = _create_new_target(level, company, name, parent_name=parent_name, kind=kind)
            if not obj:
                errors.append({'row': i + 2, 'error': 'ایجاد رکورد جدید ناموفق بود'})
                continue
        else:
            ignored += 1
            continue

        fields = _assign_target_fields(level, obj)
        is_new = action == 'new'
        entry = MappingEntry.objects.filter(
            company=company,
            source_id=source_id or None,
            level=level,
            kind=kind,
            source_code=source_code,
        ).first()
        if entry:
            entry.source_name = source_name
            entry.status = 'new' if is_new else 'matched'
            entry.resolved = True
            for k, v in fields.items():
                setattr(entry, k, v)
            entry.save()
        else:
            MappingEntry.objects.create(
                company=company,
                source_id=source_id or None,
                level=level,
                kind=kind,
                source_code=source_code,
                source_name=source_name,
                status='new' if is_new else 'matched',
                resolved=True,
                **fields,
            )
        if is_new:
            created += 1
        else:
            mapped += 1

    return Response({
        'message': 'نگاشت اعمال شد',
        'created': created,
        'mapped': mapped,
        'ignored': ignored,
        'errors': errors,
    })


def _create_new_target(level, company, name, parent_name='', kind=''):
    """ایجاد ردیف جدید در سطح کدینگ مربوطه با کد خودکار (بر اساس CodingConfig).

    `parent_name` نام/کد والد (نوع حساب برای گروه، گروه برای کل، کل برای معین) است
    که از اکسل خوانده می‌شود تا حساب در دستهٔ درست قرار گیرد.
    """
    from accounting.models import AccountType, AccountGroup
    from django.db import IntegrityError

    acc_type = AccountType.objects.filter(company=company).first()

    if level == 'group':
        # والد گروه = نوع حساب (کد یا نام)
        if parent_name:
            match = AccountType.objects.filter(company=company, code=parent_name).first()
            if not match:
                match = _best_named(AccountType.objects.filter(company=company), parent_name)
            if match:
                acc_type = match
        if not acc_type:
            return None
        for _ in range(3):
            code = suggest_code(company, 'group') or _next_code(AccountGroup, company)
            try:
                return AccountGroup.objects.create(company=company, account_type=acc_type, code=code, name=name)
            except IntegrityError:
                continue
        return None

    if level in ('general', 'subsidiary'):
        if not acc_type:
            return None
        group = None
        if parent_name:
            group = AccountGroup.objects.filter(company=company, code=parent_name).first()
            if not group:
                group = _best_named(AccountGroup.objects.filter(company=company), parent_name)
            if group:
                acc_type = group.account_type

        if level == 'general':
            for attempt in range(5):
                code = suggest_code(company, 'general') if attempt == 0 else _next_code(Account, company)
                if not code:
                    code = _next_code(Account, company)
                try:
                    return Account.objects.create(
                        company=company, account_type=acc_type, code=code, name=name,
                        level=1, group=group,
                        nature=group.nature if group else acc_type.default_nature if hasattr(acc_type, 'default_nature') else 'debit',
                    )
                except IntegrityError:
                    continue
            return None

        # معین: والد = حساب کل
        parent = None
        if parent_name:
            # ابتدا بر اساس کد، سپس نام
            parent = Account.objects.filter(company=company, parent__isnull=True, code=parent_name).first()
            if not parent:
                parent = _best_named(Account.objects.filter(company=company, parent__isnull=True), parent_name)
        if not parent:
            parent = Account.objects.filter(company=company, parent__isnull=True).first()
        if not parent:
            parent = Account.objects.create(
                company=company, account_type=acc_type,
                code=suggest_code(company, 'general') or _next_code(Account, company),
                name='حساب کل عمومی', level=1,
            )
        for attempt in range(5):
            code = suggest_code(company, 'subsidiary', parent.code) if attempt == 0 else _next_code(Account, company)
            if not code:
                code = _next_code(Account, company)
            try:
                return Account.objects.create(
                    company=company, account_type=parent.account_type or acc_type, code=code, name=name,
                    level=2, parent=parent, group=parent.group,
                    nature=parent.nature,
                )
            except IntegrityError:
                continue
        return None

    if level == 'auxiliary':
        category = None
        from accounting.models import AuxiliaryCategory
        # نام دسته: از ستون «دسته» یا از kind انتخاب‌شده
        cat_name = parent_name or dict((s['kind'], s['category']) for s in AUX_SOURCES).get(kind, '')
        if cat_name:
            category = AuxiliaryCategory.objects.filter(company=company, name=cat_name).first()
            if not category:
                category = _best_named(AuxiliaryCategory.objects.filter(company=company), cat_name)
            if not category:
                category = AuxiliaryCategory.objects.create(
                    company=company, code=cat_name[:30], name=cat_name,
                )
        for attempt in range(5):
            code = suggest_code(company, 'auxiliary') if attempt == 0 else _next_code(AuxiliaryAccount, company)
            if not code:
                code = _next_code(AuxiliaryAccount, company)
            try:
                return AuxiliaryAccount.objects.create(company=company, code=code, name=name, category=category)
            except IntegrityError:
                continue
        return None
    return None


def _best_named(qs, name):
    """بهترین تطبیق نام/کد در یک queryset بر اساس شباهت."""
    from datamapping.services import similarity
    best = None
    best_score = 0.0
    for obj in qs:
        score = max(similarity(name, obj.name), similarity(name, str(obj.code)))
        if score > best_score:
            best, best_score = obj, score
    if best and best_score >= 0.5:
        return best
    return None


def _next_code(model, company):
    """تولید کد عددی بعدی (بر اساس بیشترین کد عددی موجود)."""
    from django.db.models import Max
    agg = model.objects.filter(company=company).aggregate(m=Max('code'))
    cur = 0
    for ch in str(agg.get('m') or ''):
        if ch.isdigit():
            cur = cur * 10 + int(ch)
    return str(cur + 1) if cur else '1'


# -----------------------------------------------------------------------------
# لیست نگاشت‌ها
# -----------------------------------------------------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def entries(request):
    company = _company(request)
    level = request.query_params.get('level')
    kind = request.query_params.get('kind')
    qs = MappingEntry.objects.filter(company=company).select_related(
        'source', 'target_group', 'target_account', 'target_auxiliary',
    )
    if level:
        qs = qs.filter(level=level)
    if kind:
        qs = qs.filter(kind=kind)
    return Response(MappingEntrySerializer(qs[:1000], many=True).data)


@api_view(['PATCH', 'DELETE'])
@permission_classes([IsAuthenticated])
def entry_detail(request, pk):
    """اصلاح یا حذف یک نگاشت.

    PATCH: تغییر «هدف» (target_group/target_account/target_auxiliary) یا وضعیت/نادیده‌گرفتن.
    DELETE: حذف نگاشت.
    """
    company = _company(request)
    entry = MappingEntry.objects.filter(company=company, id=pk).first()
    if not entry:
        return Response({'error': 'نگاشت یافت نشد'}, status=404)

    if request.method == 'DELETE':
        entry.delete()
        return Response({'message': 'نگاشت حذف شد'})

    data = request.data
    level = entry.level
    target_id = data.get('target_id')
    if target_id is not None:
        obj = _resolve_target(level, target_id, company)
        for k, v in _assign_target_fields(level, obj).items():
            setattr(entry, k, v)
        # پاک‌کردن هدف‌های دیگر سطح در صورت تغییر سطح (ایمن‌سازی)
        if level == 'group':
            entry.target_account = None
            entry.target_auxiliary = None
        elif level in ('general', 'subsidiary'):
            entry.target_group = None
            entry.target_auxiliary = None
        else:
            entry.target_group = None
            entry.target_account = None

    status_val = data.get('status')
    if status_val:
        entry.status = status_val
    if 'resolved' in data:
        entry.resolved = bool(data['resolved'])
    # بر اساس هدف، وضعیت را به‌روز کن
    if entry.target_group or entry.target_account or entry.target_auxiliary:
        entry.resolved = True
        if entry.status == 'pending':
            entry.status = 'matched'
    entry.save()
    return Response(MappingEntrySerializer(entry).data)


# -----------------------------------------------------------------------------
# ایمپورت اسناد حسابداری
# -----------------------------------------------------------------------------
DOC_HEADERS = {
    'row_no': ['ردیف', 'ردیف سند', 'سطر', 'row', 'row_no'],
    'date': ['تاریخ سند', 'تاریخ', 'date'],
    'doc_number': ['شماره سند', 'شماره', 'شماره سند مبدا', 'number', 'doc_no'],
    'reference': ['شماره عطف', 'عطف', 'reference', 'ref'],
    'account_code': ['معین', 'کد معین', 'کد حساب', 'کد', 'حساب', 'account', 'account_code'],
    'auxiliary_1': ['تفصیل1', 'تفصیل 1', 'تفصیل یک', 'تفصیلی1', 'تفصیلی یک', 'aux1', 'auxiliary1'],
    'auxiliary_2': ['تفصیل2', 'تفصیل 2', 'تفصیل دو', 'تفصیلی2', 'تفصیلی دو', 'aux2', 'auxiliary2'],
    'auxiliary_3': ['تفصیل3', 'تفصیل 3', 'تفصیل سه', 'تفصیلی3', 'تفصیلی سه', 'aux3', 'auxiliary3'],
    'line_description': ['شرح سطر', 'شرح ردیف', 'شرح', 'line_description'],
    'debit': ['بدهکار', 'debit'],
    'credit': ['بستانکار', 'credit'],
    'doc_description': ['شرح سند', 'شرح کل سند', 'doc_description'],
}


def _parse_date(value):
    """تبدیل تاریخ به میلادی. ورودی ممکن است شمسی متنی، میلادی متنی یا سلول تاریخ اکسل باشد.

    - اگر مقدار سلول اکسل از نوع date/datetime باشد، مستقیم میلادی برمی‌گردد (openpyxl
      تاریخ را میلادی می‌دهد) و نیازی به تبدیل ندارد.
    - اگر رشتهٔ متنی باشد: سال 1300 تا 1499 شمسی فرض شده و با jdatetime به میلادی تبدیل
      می‌شود؛ در غیر این صورت میلادی تفسیر می‌شود.
    """
    import jdatetime

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    s = str(value).strip()
    if not s:
        return None

    # نرمال‌سازی ارقام فارسی/عربی و جداکننده‌ها
    s = normalize(s).replace('/', '-').replace('.', '-').replace('،', '-')

    # فرمت شمسی: 1403/05/12 یا 1403-05-12
    for fmt in ('%Y-%m-%d', '%Y/%m/%d'):
        try:
            d = datetime.strptime(s, fmt)
        except ValueError:
            continue
        if 1300 <= d.year <= 1499:
            try:
                return jdatetime.date(d.year, d.month, d.day).togregorian()
            except Exception:
                return None
        return d.date()

    # تاریخ هجری شمسی بدون جداکننده: 14030512
    if s.isdigit() and len(s) in (8,):
        try:
            y, m, d = int(s[0:4]), int(s[4:6]), int(s[6:8])
            if 1300 <= y <= 1499:
                return jdatetime.date(y, m, d).togregorian()
        except Exception:
            pass
    return None


def _resolve_account(company, code, source_id=None):
    if not code:
        return None
    acc = Account.objects.filter(company=company, code=code).first()
    if acc:
        return acc
    # ابتدا نگاشت مختص منبع، سپس نگاشت عمومی
    qs = MappingEntry.objects.filter(
        company=company, level__in=['subsidiary', 'general'], source_code=code, resolved=True,
    )
    if source_id:
        entry = qs.filter(source_id=source_id).first() or qs.filter(source__isnull=True).first()
    else:
        entry = qs.first()
    if entry and entry.target_account:
        return entry.target_account
    return None


def _resolve_auxiliary(company, code, source_id=None):
    if not code:
        return None
    aux = AuxiliaryAccount.objects.filter(company=company, code=code).first()
    if aux:
        return aux
    qs = MappingEntry.objects.filter(
        company=company, level='auxiliary', source_code=code, resolved=True,
    )
    if source_id:
        entry = qs.filter(source_id=source_id).first() or qs.filter(source__isnull=True).first()
    else:
        entry = qs.first()
    if entry and entry.target_auxiliary:
        return entry.target_auxiliary
    return None


def _fiscal_year_for(company, doc_date, fiscal_year_id=None):
    """سال مالی: یا از id ارسالی، یا با تشخیص از بازهٔ تاریخ."""
    if fiscal_year_id:
        fy = FiscalYear.objects.filter(company=company, id=fiscal_year_id).first()
        if fy:
            return fy
    if doc_date:
        fy = FiscalYear.objects.filter(
            company=company, start_date__lte=doc_date, end_date__gte=doc_date,
        ).order_by('-start_date').first()
        if fy:
            return fy
    return FiscalYear.objects.filter(company=company).order_by('-start_date').first()


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def import_documents(request):
    company = _company(request)
    post_now = str(request.data.get('post') or 'false').lower() in ('true', '1')
    source_id = request.data.get('source_id') or None
    fiscal_year_id = request.data.get('fiscal_year') or request.query_params.get('fiscal_year') or None
    file = request.FILES.get('file')
    if not file:
        return Response({'error': 'فایل انتخاب نشده است'}, status=400)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        raw = list(ws.iter_rows(values_only=True))
        if not raw:
            return Response({'error': 'فایل خالی است'}, status=400)
        headers = [str(h).strip() if h is not None else '' for h in raw[0]]
        col_map = _match_headers(headers, DOC_HEADERS)
        if 'account_code' not in col_map:
            return Response({'error': 'ستون «معین / کد حساب» پیدا نشد'}, status=400)

        def _text(v):
            return str(v).strip() if v is not None else ''

        rows = []
        for r in raw[1:]:
            if not any(r):
                continue

            def cell(key):
                if key in col_map and col_map[key] < len(r):
                    return r[col_map[key]]
                return None

            rows.append({
                'row_no': _text(cell('row_no')),
                'doc_number': _text(cell('doc_number')),
                'reference': _text(cell('reference')),
                'date': _parse_date(cell('date')),
                'doc_description': _text(cell('doc_description')),
                'account_code': _text(cell('account_code')),
                'auxiliary_1': _text(cell('auxiliary_1')),
                'auxiliary_2': _text(cell('auxiliary_2')),
                'auxiliary_3': _text(cell('auxiliary_3')),
                'line_description': _text(cell('line_description')),
                'debit': to_decimal(cell('debit')),
                'credit': to_decimal(cell('credit')),
            })

        from collections import OrderedDict
        docs = OrderedDict()
        for r in rows:
            key = r['doc_number'] or (r['reference'] or (r['date'].isoformat() if r['date'] else 'doc'))
            docs.setdefault(key, []).append(r)

        journal, _ = Journal.objects.get_or_create(
            company=company, code='GEN',
            defaults={'name': 'روزنامه عمومی', 'journal_type': 'manual'},
        )

        created_docs = 0
        created_lines = 0
        unresolved_codes = set()
        skipped_rows = []
        for doc_key, lines in docs.items():
            first = lines[0]
            doc_date = first['date'] or date.today()
            fiscal_year = _fiscal_year_for(company, doc_date, fiscal_year_id)

            doc = AccountingDocument.objects.create(
                company=company,
                journal=journal,
                fiscal_year=fiscal_year,
                date=doc_date,
                description=first['doc_description'] or f'سند ایمپورت‌شده {doc_key}',
                reference=first['reference'] or '',
                number='',
                status='posted' if post_now else 'draft',
                source_module='datamapping',
                source_type='excel',
                source_id=str(doc_key),
                created_by=request.user if request.user.is_authenticated else None,
            )
            created_docs += 1
            for i, ln in enumerate(lines, start=1):
                acc = _resolve_account(company, ln['account_code'], source_id)
                if not acc:
                    if ln['account_code']:
                        unresolved_codes.add(ln['account_code'])
                    skipped_rows.append({'row': ln['row_no'] or i, 'reason': 'کد معین یافت نشد', 'code': ln['account_code']})
                    continue
                aux1 = _resolve_auxiliary(company, ln['auxiliary_1'], source_id)
                aux2 = _resolve_auxiliary(company, ln['auxiliary_2'], source_id)
                aux3 = _resolve_auxiliary(company, ln['auxiliary_3'], source_id)
                for code, aux in ((ln['auxiliary_1'], aux1), (ln['auxiliary_2'], aux2), (ln['auxiliary_3'], aux3)):
                    if code and not aux:
                        unresolved_codes.add(code)

                AccountingDocumentLine.objects.create(
                    company=company,
                    document=doc,
                    account=acc,
                    auxiliary_1=aux1,
                    auxiliary_2=aux2,
                    auxiliary_3=aux3,
                    line_no=i,
                    description=ln['line_description'] or '',
                    debit=ln['debit'],
                    credit=ln['credit'],
                )
                created_lines += 1
                _link_aux_categories(acc, aux1, aux2, aux3)

        return Response({
            'message': 'ایمپورت اسناد انجام شد',
            'created_documents': created_docs,
            'created_lines': created_lines,
            'post_now': post_now,
            'unresolved_codes': sorted(unresolved_codes)[:100],
            'skipped_rows': skipped_rows[:100],
        })
    except Exception as e:
        return Response({'error': f'خطا در ایمپورت: {str(e)[:150]}'}, status=400)


def _link_aux_categories(account, aux1, aux2, aux3):
    """ایجاد رابطهٔ دستهٔ تفصیلی با حساب معین.

    وقتی سندی با یک معین و تفصیل‌هایی از یک دسته ثبت می‌شود، حساب معین باید آن
    دسته را به‌عنوان `auxiliary_category_1/2/3` بشناسد تا در سند بعدی همان
    تفصیل‌ها قابل انتخاب و درج کد باشند.
    """
    changed = False
    for idx, aux in enumerate((aux1, aux2, aux3), start=1):
        cat = getattr(aux, 'category', None)
        if not cat:
            continue
        field = f'auxiliary_category_{idx}'
        if not getattr(account, field, None):
            setattr(account, field, cat)
            changed = True
    if changed:
        account.save(update_fields=[
            'auxiliary_category_1', 'auxiliary_category_2', 'auxiliary_category_3', 'updated_at',
        ])


# -----------------------------------------------------------------------------
# ساخت خودکار حساب‌های تفصیلی از موجودیت‌های ماژول‌های دیگر (دسته به دسته)
# -----------------------------------------------------------------------------
AUX_SOURCES = [
    {'kind': 'bank', 'label': 'بانک‌ها و صندوق‌ها', 'module': 'خزانه‌داری', 'category': 'بانک', 'cat_source': 'bank'},
    {'kind': 'employee', 'label': 'پرسنل', 'module': 'منابع انسانی', 'category': 'پرسنل', 'cat_source': 'employee'},
    {'kind': 'supplier', 'label': 'تأمین‌کنندگان', 'module': 'خرید و تدارکات', 'category': 'تأمین‌کننده', 'cat_source': 'manual'},
    {'kind': 'custodian', 'label': 'تنخواه‌داران', 'module': 'تنخواه', 'category': 'تنخواه‌دار', 'cat_source': 'employee'},
    {'kind': 'party', 'label': 'طرف‌حساب‌ها (حقوقی/حقیقی)', 'module': 'قراردادها', 'category': 'طرف حساب', 'cat_source': 'party'},
    {'kind': 'contract', 'label': 'قراردادها', 'module': 'قراردادها', 'category': 'قرارداد', 'cat_source': 'contract'},
    {'kind': 'project', 'label': 'پروژه‌ها', 'module': 'پروژه‌ها', 'category': 'پروژه', 'cat_source': 'project'},
]


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def auxiliary_sources(request):
    """لیست منابع موجود برای ساخت تفصیلی + تعداد هر کدام."""
    company = _company(request)

    from contracts.models import ContractParty, Contract
    from treasury.models import TreasuryEntity
    from projects.models import Project
    from employees.models import Employee
    from procurement.models import Supplier
    from pettycash.models import PettyCashFund

    counts = {
        'bank': TreasuryEntity.objects.filter(company=company, is_active=True).count(),
        'employee': Employee.objects.filter(company=company, is_active=True).count(),
        'supplier': Supplier.objects.filter(company=company, status='active').count(),
        'custodian': PettyCashFund.objects.filter(company=company, status='active').values('custodian_id').distinct().count(),
        'party': ContractParty.objects.filter(company=company, is_active=True).count(),
        'contract': Contract.objects.filter(company=company).count(),
        'project': Project.objects.filter(company=company).count(),
    }

    return Response([
        {
            'kind': s['kind'],
            'label': s['label'],
            'module': s['module'],
            'category': s['category'],
            'count': counts.get(s['kind'], 0),
        }
        for s in AUX_SOURCES
    ])


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def sync_auxiliaries(request):
    """ساخت تفصیل‌ها از ماژول‌ها. `kinds` لیست منابع؛ اگر خالی بود همه را همگام می‌کند."""
    company = _company(request)
    kinds = request.data.get('kinds') or []
    if isinstance(kinds, str):
        kinds = [kinds]
    kinds = [k for k in kinds if any(s['kind'] == k for s in AUX_SOURCES)]
    if not kinds:
        kinds = [s['kind'] for s in AUX_SOURCES]

    from accounting.models import AuxiliaryCategory
    from contracts.models import ContractParty, Contract
    from treasury.models import TreasuryEntity
    from projects.models import Project
    from employees.models import Employee
    from procurement.models import Supplier
    from pettycash.models import PettyCashFund

    def get_category(name, source, sort):
        return AuxiliaryCategory.objects.get_or_create(
            company=company, name=name,
            defaults={'code': source + str(sort), 'source': source, 'sort_order': sort},
        )[0]

    cat_bank = get_category('بانک', 'bank', 1)
    cat_employee = get_category('پرسنل', 'employee', 2)
    cat_supplier = get_category('تأمین‌کننده', 'manual', 3)
    cat_custodian = get_category('تنخواه‌دار', 'employee', 4)
    cat_party_legal = get_category('طرف حساب حقوقی', 'party', 5)
    cat_party_natural = get_category('طرف حساب حقیقی', 'party', 6)
    cat_contract = get_category('قرارداد', 'contract', 7)
    cat_project = get_category('پروژه', 'project', 8)

    created = 0
    updated = 0
    counts = {}

    def upsert(lookup, aux_type, category, name, extra=None):
        nonlocal created, updated
        obj = AuxiliaryAccount.objects.filter(company=company, **lookup).first()
        if obj:
            changed = False
            if not obj.category:
                obj.category = category
                changed = True
            if not obj.aux_type or obj.aux_type == 'other':
                obj.aux_type = aux_type
                changed = True
            if changed:
                obj.save(update_fields=['category', 'aux_type', 'updated_at'])
            updated += 1
            return
        kwargs = {'company': company, 'category': category, 'aux_type': aux_type, 'name': name, **lookup}
        if extra:
            kwargs.update(extra)
        kwargs['code'] = suggest_code(company, 'auxiliary') or _next_code(AuxiliaryAccount, company)
        AuxiliaryAccount.objects.create(**kwargs)
        created += 1

    if 'bank' in kinds:
        for entity in TreasuryEntity.objects.filter(company=company, is_active=True):
            upsert({'treasury_entity': entity}, 'bank' if entity.entity_type == 'bank' else 'cash', cat_bank, entity.name)
        counts['bank'] = TreasuryEntity.objects.filter(company=company, is_active=True).count()

    if 'employee' in kinds:
        for emp in Employee.objects.filter(company=company, is_active=True):
            upsert({'employee': emp}, 'employee', cat_employee, emp.full_name)
        counts['employee'] = Employee.objects.filter(company=company, is_active=True).count()

    if 'supplier' in kinds:
        for sup in Supplier.objects.filter(company=company, status='active'):
            upsert({'supplier': sup}, 'supplier', cat_supplier, sup.name)
        counts['supplier'] = Supplier.objects.filter(company=company, status='active').count()

    if 'custodian' in kinds:
        seen = set()
        for fund in PettyCashFund.objects.filter(company=company, status='active').select_related('custodian'):
            emp = fund.custodian
            if not emp or emp.id in seen:
                continue
            seen.add(emp.id)
            upsert({'employee': emp}, 'employee', cat_custodian, emp.full_name)
        counts['custodian'] = len(seen)

    if 'party' in kinds:
        for party in ContractParty.objects.filter(company=company, is_active=True):
            cat = cat_party_legal if party.person_type == 'legal' else cat_party_natural
            upsert({'party': party}, 'party', cat, party.name, {'person_type': party.person_type})
        counts['party'] = ContractParty.objects.filter(company=company, is_active=True).count()

    if 'contract' in kinds:
        for contract in Contract.objects.filter(company=company).select_related('party'):
            if contract.party:
                label = f"{contract.party.name} - {contract.number}".strip(' -')
            else:
                label = contract.number or f'قرارداد {contract.pk}'
            upsert({'contract': contract}, 'contract', cat_contract, label, {'party': contract.party})
        counts['contract'] = Contract.objects.filter(company=company).count()

    if 'project' in kinds:
        for project in Project.objects.filter(company=company):
            upsert({'project': project}, 'project', cat_project, project.name)
        counts['project'] = Project.objects.filter(company=company).count()

    return Response({
        'message': 'تفصیل‌ها همگام‌سازی شدند',
        'created': created,
        'updated': updated,
        'counts': counts,
    })


# -----------------------------------------------------------------------------
# قالب اکسل نمونه
# -----------------------------------------------------------------------------
TEMPLATE_HEADERS = {
    'group': ['کد', 'عنوان', 'کد والد'],
    'general': ['کد کل', 'عنوان', 'کد والد'],
    'subsidiary': ['کد حساب', 'عنوان', 'کد والد'],
    'auxiliary': ['کد', 'عنوان', 'دسته'],
    'documents': ['ردیف', 'تاریخ سند', 'شماره سند', 'شماره عطف', 'معین', 'تفصیل1', 'تفصیل2', 'تفصیل3', 'شرح سطر', 'بدهکار', 'بستانکار', 'شرح سند'],
}
TEMPLATE_SAMPLES = {
    'group': ['100', 'دارایی‌های جاری', '1'],
    'general': ['1001', 'موجودی نقد', '100'],
    'subsidiary': ['10011', 'صندوق', '1001'],
    'auxiliary': ['100', 'بانک ملت', 'بانک'],
    'documents': ['1', '1403/12/25', '1001', 'عطف-۱', '10011', '100', '', '', 'خرید نقدی کالا', '15000000', '', 'سند خرید'],
}


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def template(request):
    """دانلود فایل اکسل نمونه برای یک سطح کدینگ."""
    level = request.query_params.get('level', 'subsidiary')
    headers = TEMPLATE_HEADERS.get(level, TEMPLATE_HEADERS['subsidiary'])
    sample = TEMPLATE_SAMPLES.get(level, TEMPLATE_SAMPLES['subsidiary'])

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'اسناد حسابداری' if level == 'documents' else LEVEL_LABELS.get(level, 'کدینگ')

    header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    center = Alignment(horizontal='center', vertical='center')

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    for col_idx, val in enumerate(sample, start=1):
        ws.cell(row=2, column=col_idx, value=val).alignment = center

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=template_{level}.xlsx'
    wb.save(response)
    return response


# -----------------------------------------------------------------------------
# حذف همهٔ کدینگ‌های حسابداری (برای شروع مجدد ایمپورت)
# -----------------------------------------------------------------------------
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def clear_codings(request):
    """حذف همهٔ گروه‌ها، حساب‌ها (کل/معین)، تفصیلی‌ها، نگاشت‌ها و اسناد حسابداری شرکت.

    برای شروع مجدد فرآیند ایمپورت. اسناد و خطوط هم حذف می‌شوند چون با PROTECT
    به حساب‌ها متصل هستند. CodingConfig (سربرگ کدینگ) دست‌نخورده می‌ماند.
    """
    company = _company(request)
    confirm = str(request.data.get('confirm') or '').strip()
    if confirm != 'DELETE':
        return Response({'error': 'برای حذف، مقدار confirm را "DELETE" بفرستید.'}, status=400)

    from accounting.models import (
        AccountGroup, Account, AuxiliaryAccount,
        AccountingDocument, AccountingDocumentLine,
    )

    result = {}

    def _count(qs):
        return qs.count()

    # 1) خطوط و اسناد (به‌خاطر PROTECT حساب)
    result['document_lines'] = _count(AccountingDocumentLine.objects.filter(company=company))
    AccountingDocumentLine.objects.filter(company=company).delete()
    result['documents'] = _count(AccountingDocument.objects.filter(company=company))
    AccountingDocument.objects.filter(company=company).delete()

    # 2) تفصیلی‌ها
    result['auxiliaries'] = _count(AuxiliaryAccount.objects.filter(company=company))
    AuxiliaryAccount.objects.filter(company=company).delete()

    # 3) حساب‌های معین (دارای parent) سپس کل (بدون parent)
    result['subsidiaries'] = _count(Account.objects.filter(company=company, parent__isnull=False))
    Account.objects.filter(company=company, parent__isnull=False).delete()
    result['generals'] = _count(Account.objects.filter(company=company, parent__isnull=True))
    Account.objects.filter(company=company, parent__isnull=True).delete()

    # 4) گروه‌ها
    result['groups'] = _count(AccountGroup.objects.filter(company=company))
    AccountGroup.objects.filter(company=company).delete()

    # 5) نگاشت‌ها
    result['mappings'] = _count(MappingEntry.objects.filter(company=company))
    MappingEntry.objects.filter(company=company).delete()

    return Response({'message': 'همهٔ کدینگ‌ها حذف شدند', 'deleted': result})
