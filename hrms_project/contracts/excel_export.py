"""خروجی اکسل قراردادها — تک‌شیت و **یک ردیف برای هر قرارداد**.

همهٔ اقلام درون هر قرارداد (فاکتورها، صورت‌وضعیت‌ها و اضافه/کسورهای آن، پرداخت‌ها،
تضمین‌ها، الحاقیه‌ها و بندهای آن، تخصیص پروژه‌ها، پیش‌پرداخت‌های مرحله‌ای، اقلام ضمانت و
اختلافات) به‌صورت ستون‌های شماره‌دار در همان ردیف قرارداد می‌آیند؛ مثلاً «فاکتور ۱ — مبلغ»،
«پرداخت ۲ — تاریخ»، «تضمین ۱ — شماره». تعداد این ستون‌ها از روی بیشترین تعداد آیتم در بین
قراردادهای خروجی ساخته می‌شود (بدون محدودیت از پیش تعیین‌شده).
"""
from decimal import Decimal

from contracts.numbering import display_statement_number

try:  # تاریخ شمسی (همان الگوی ماژول پروژه‌ها)
    from jdatetime import date as jdate
except ImportError:  # pragma: no cover
    jdate = None


# ----------------------------------------------------------------------------
# کمک‌تابع‌های ایمن (نبود هیچ فیلد/رابطه‌ای نباید خروجی را قطع کند)
# ----------------------------------------------------------------------------
def _fa_date(value):
    """تاریخ شمسی به شکل 1404/01/15."""
    if not value:
        return ''
    if jdate is None:
        return str(value)
    try:
        j = jdate.fromgregorian(date=value)
        return f'{j.year:04d}/{j.month:02d}/{j.day:02d}'
    except Exception:  # noqa: BLE001
        return str(value)


def _num(value):
    """عدد (مبلغ/درصد) یا رشتهٔ خالی."""
    if value in (None, ''):
        return ''
    try:
        return float(Decimal(value))
    except Exception:  # noqa: BLE001
        return ''


def _text(value):
    return '' if value in (None, '') else str(value)


def _attr(obj, name):
    """مقدار یک فیلد به‌صورت متن (نبود فیلد مشکلی ایجاد نمی‌کند)."""
    return _text(getattr(obj, name, ''))


def _related_name(obj, field, *attrs):
    """نام یک رابطه (FK) را ایمن بخوان (نام فیلد در مدل‌ها متفاوت است)."""
    related = getattr(obj, field, None)
    if related is None:
        return ''
    for attr in attrs:
        value = getattr(related, attr, None)
        if value:
            return str(value)
    return ''


def _display(obj, attr):
    """متن نمایشی یک فیلد دارای choices (get_<attr>_display)."""
    if obj is None:
        return ''
    getter = getattr(obj, f'get_{attr}_display', None)
    if callable(getter):
        try:
            return getter()
        except Exception:  # noqa: BLE001
            return ''
    return ''


def _party_name(contract):
    return _related_name(contract, 'party', 'name')


def _party_type(contract):
    party = getattr(contract, 'party', None)
    if party is None or not getattr(contract, 'party_id', None):
        return ''
    return _display(party, 'person_type')


def _type_label(contract):
    return _related_name(contract, 'contract_type_master', 'name') or _display(contract, 'contract_type')


def _contract_currency(contract):
    return _related_name(contract, 'currency', 'name', 'code') or 'ریال'


def _attr_list(contract, field):
    """فهرست JSON یک قرارداد (همیشه list برگردانده می‌شود)."""
    value = getattr(contract, field, None)
    return value if isinstance(value, list) else []


def _related_list(contract, field):
    """فهرست آیتم‌های یک رابطه (اگر در دسترس نباشد، فهرست خالی)."""
    fetcher = getattr(getattr(contract, field, None), 'all', None)
    if not callable(fetcher):
        return []
    try:
        return list(fetcher())
    except Exception:  # noqa: BLE001
        return []


def _related_count(contract, field):
    """تعداد اسناد/روابط مرتبط (اگر در دسترس نباشد، رشتهٔ خالی)."""
    counter = getattr(getattr(contract, field, None), 'count', None)
    if not callable(counter):
        return ''
    try:
        return counter()
    except Exception:  # noqa: BLE001
        return ''


def _project_label(project_id):
    """نام یک پروژه با شناسهٔ عددی (ایمن در برابر نبود جدول پروژه)."""
    if not project_id:
        return ''
    try:
        from projects.models import Project
        name = Project.objects.filter(id=project_id).values_list('name', flat=True).first()
        return name or f'#{project_id}'
    except Exception:  # noqa: BLE001
        return f'#{project_id}'


def _project_names(contract):
    """نام پروژه‌های تخصیص‌یافته (از project_allocations یا پروژهٔ منفرد)."""
    allocations = _attr_list(contract, 'project_allocations')
    ids = [item.get('project') for item in allocations if isinstance(item, dict) and item.get('project')]
    if getattr(contract, 'project_id', None) and not ids:
        ids.append(contract.project_id)
    if not ids:
        return ''
    return '، '.join(_project_label(i) for i in ids)


def _join(label, value):
    """ساخت بخشی از یک ستون جزئیات."""
    text = _text(value)
    return f'{label}: {text}' if text else ''


def _details(*parts):
    return '؛ '.join(p for p in parts if p)


def _sort_key(obj):
    """کلید مرتب‌سازی ایمن بر اساس اولین تاریخ موجود."""
    day = (getattr(obj, 'date', None) or getattr(obj, 'issue_date', None)
           or getattr(obj, 'opened_date', None))
    return (str(day or ''), getattr(obj, 'id', 0) or 0)


def _day_of(value):
    """تاریخ (بدون ساعت) از یک تاریخ/تاریخ‌زمان."""
    return getattr(value, 'date', lambda: None)() if value is not None else None


def _iso_day(value):
    """رشتهٔ ISO تاریخ (خروجی موتور الحاقیه) را به date تبدیل می‌کند."""
    if not value:
        return None
    if not isinstance(value, str):
        return value
    from datetime import date as _date

    try:
        return _date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _created_day(contract):
    return _day_of(getattr(contract, 'created_at', None))


def _updated_day(contract):
    return _day_of(getattr(contract, 'updated_at', None))


def _sum_of(items, *fields):
    """جمع اولین فیلد موجود از میان چند نام ممکن."""
    total = 0.0
    for item in items:
        for field in fields:
            value = getattr(item, field, None)
            if value not in (None, ''):
                total += float(value or 0)
                break
    return total


# ----------------------------------------------------------------------------
# «برگه» هر قرارداد: فهرست آیتم‌های هر نوع + جمع‌های مالی
# ----------------------------------------------------------------------------
def _bundle(contract):
    """آیتم‌های درون یک قرارداد + جمع‌های مالی آن."""
    invoices = sorted(_related_list(contract, 'invoices'), key=_sort_key)
    statements = sorted(_related_list(contract, 'statements'), key=_sort_key)
    payments = sorted(_related_list(contract, 'payments'), key=_sort_key)
    guarantees = sorted(_related_list(contract, 'guarantees'), key=_sort_key, reverse=True)
    addendums = sorted(_related_list(contract, 'addendums'), key=_sort_key)
    disputes = sorted(_related_list(contract, 'disputes'), key=_sort_key, reverse=True)
    payments_by_invoice, payments_by_statement = {}, {}
    for payment in payments:
        invoice_id = getattr(payment, 'invoice_id', None)
        statement_id = getattr(payment, 'statement_id', None)
        if invoice_id:
            payments_by_invoice.setdefault(invoice_id, []).append(payment)
        if statement_id:
            payments_by_statement.setdefault(statement_id, []).append(payment)
    return {
        'contract': contract,
        'invoices': invoices,
        'statements': statements,
        'payments': payments,
        'guarantees': guarantees,
        'addendums': addendums,
        'disputes': disputes,
        'payments_by_invoice': payments_by_invoice,
        'payments_by_statement': payments_by_statement,
        'invoice_numbers': {getattr(i, 'id', None): getattr(i, 'number', '') for i in invoices},
        'statement_numbers': {getattr(s, 'id', None): display_statement_number(getattr(s, 'number', '')) for s in statements},
        'invoice_total': _sum_of(invoices, 'total', 'amount'),
        'statement_total': _sum_of(statements, 'net_amount', 'amount'),
        'payment_total': _sum_of(payments, 'amount'),
        'guarantee_total': _sum_of(guarantees, 'amount'),
        'addendum_total': _sum_of(addendums, 'amount_change'),
        'dispute_total': _sum_of(disputes, 'claim_amount'),
    }


def _invoice_slots(bundle):
    """ستون‌های «فاکتور N»."""
    slots = []
    for invoice in bundle['invoices']:
        linked = bundle['payments_by_invoice'].get(getattr(invoice, 'id', None), [])
        slots.append({
            'number': _attr(invoice, 'number'),
            'subject': _attr(invoice, 'subject'),
            'date': _fa_date(getattr(invoice, 'date', None)),
            'due_date': _fa_date(getattr(invoice, 'due_date', None)),
            'amount': _num(getattr(invoice, 'amount', None)),
            'vat': _num(getattr(invoice, 'vat', None)),
            'total': _num(getattr(invoice, 'total', None)),
            'status': 'پرداخت‌شده' if getattr(invoice, 'is_paid', False) else 'پرداخت‌نشده',
            'payments': f"{len(linked)} پرداخت / {sum(float(getattr(p, 'amount', None) or 0) for p in linked):,.0f}",
            'description': _attr(invoice, 'description'),
        })
    return slots


def _statement_slots(bundle):
    """ستون‌های «صورت‌وضعیت N»."""
    slots = []
    for statement in bundle['statements']:
        linked = bundle['payments_by_statement'].get(getattr(statement, 'id', None), [])
        slots.append({
            'number': display_statement_number(getattr(statement, 'number', '')),
            'date': _fa_date(getattr(statement, 'date', None)),
            'cumulative': _num(getattr(statement, 'cumulative_previous_amount', None)),
            'work_done': _num(getattr(statement, 'work_done', None)),
            'amount': _num(getattr(statement, 'amount', None)),
            'vat': _num(getattr(statement, 'value_added_tax', None)),
            'additions_total': _num(getattr(statement, 'additions_total', None)),
            'deductions_total': _num(getattr(statement, 'deductions_total', None)),
            'net_amount': _num(getattr(statement, 'net_amount', None)),
            'status': 'تأییدشده' if getattr(statement, 'is_approved', False) else 'تأییدنشده',
            'payments': f"{len(linked)} پرداخت / {sum(float(getattr(p, 'amount', None) or 0) for p in linked):,.0f}",
            'description': _attr(statement, 'description'),
        })
    return slots


def _statement_extra_slots(bundle, field, prefix):
    """ستون‌های «اضافهٔ صورت‌وضعیت N» یا «کسور صورت‌وضعیت N»."""
    slots = []
    for statement in bundle['statements']:
        for item in (getattr(statement, field, None) or []):
            if not isinstance(item, dict):
                continue
            slots.append({
                'statement': display_statement_number(getattr(statement, 'number', '')),
                'title': f"{prefix}{item.get('title') or item.get('description') or ''}".strip(),
                'amount': _num(item.get('amount')),
                'percent': _num(item.get('percent')),
                'note': _text(item.get('note')),
            })
    return slots


def _payment_slots(bundle):
    """ستون‌های «پرداخت N»."""
    slots = []
    for payment in bundle['payments']:
        invoice_id = getattr(payment, 'invoice_id', None)
        statement_id = getattr(payment, 'statement_id', None)
        slots.append({
            'reference': _attr(payment, 'reference'),
            'date': _fa_date(getattr(payment, 'date', None)),
            'amount': _num(getattr(payment, 'amount', None)),
            'method': _attr(payment, 'method'),
            'linked': _details(
                _join('فاکتور', bundle['invoice_numbers'].get(invoice_id, invoice_id or '')),
                _join('صورت‌وضعیت', bundle['statement_numbers'].get(statement_id, statement_id or '')),
            ),
            'note': _attr(payment, 'note'),
        })
    return slots


def _guarantee_slots(bundle):
    """ستون‌های «تضمین N»."""
    slots = []
    for guarantee in bundle['guarantees']:
        slots.append({
            'number': _attr(guarantee, 'number'),
            'type': _display(guarantee, 'guarantee_type'),
            'instrument': _display(guarantee, 'instrument_type'),
            'issue_date': _fa_date(getattr(guarantee, 'issue_date', None)),
            'expiry_date': _fa_date(getattr(guarantee, 'expiry_date', None)),
            'amount': _num(getattr(guarantee, 'amount', None)),
            'bank': _attr(guarantee, 'bank'),
            'status': _details(
                'آزادشده' if getattr(guarantee, 'is_released', False) else 'آزاد نشده',
                _join('آخرین اقدام', _display(guarantee, 'last_action')),
                _join('تاریخ آزادسازی', _fa_date(getattr(guarantee, 'release_date', None))),
            ),
            'details': _details(
                _join('شماره چک', _attr(guarantee, 'check_number')),
                _join('بانک چک', _attr(guarantee, 'check_bank')),
                _join('سررسید چک', _fa_date(getattr(guarantee, 'check_due_date', None))),
                _join('شماره سفته', _attr(guarantee, 'promissory_number')),
                _join('سررسید سفته', _fa_date(getattr(guarantee, 'promissory_due_date', None))),
                _join('شماره ضمانت‌نامه', _attr(guarantee, 'guarantee_number')),
                _join('انقضای ضمانت‌نامه', _fa_date(getattr(guarantee, 'guarantee_expiry_date', None))),
            ),
            'note': _attr(guarantee, 'note'),
        })
    return slots


def _addendum_slots(bundle):
    """ستون‌های «الحاقیه N»."""
    slots = []
    for addendum in bundle['addendums']:
        subject_labels = getattr(addendum, 'subject_labels', None)
        if subject_labels is None:
            subject_labels = [_display(addendum, 'subject_type') or _attr(addendum, 'subject')]
        slots.append({
            'number': _attr(addendum, 'number'),
            'subjects': '، '.join(label for label in subject_labels if label),
            'subject_type': _display(addendum, 'subject_type') or _attr(addendum, 'subject'),
            'subject': _attr(addendum, 'subject'),
            'date': _fa_date(getattr(addendum, 'date', None)),
            'new_end_date': _fa_date(getattr(addendum, 'new_end_date', None)),
            'amount_change': _num(getattr(addendum, 'amount_change', None)),
            'percent_change': _num(getattr(addendum, 'percent_change', None)),
            'provisions': len(getattr(addendum, 'provisions', None) or []),
            'description': _attr(addendum, 'change_description'),
        })
    return slots


def _provision_slots(bundle):
    """ستون‌های «مفاد الحاقیه N»."""
    slots = []
    for addendum in bundle['addendums']:
        for position, provision in enumerate(getattr(addendum, 'provisions', None) or [], start=1):
            if not isinstance(provision, dict):
                continue
            slots.append({
                'addendum': _details(_attr(addendum, 'number'), f'بند {position}'),
                'title': _text(provision.get('title')),
                'text': _text(provision.get('text')),
            })
    return slots


def _allocation_slots(bundle):
    """ستون‌های «تخصیص پروژه N»."""
    slots = []
    for item in _attr_list(bundle['contract'], 'project_allocations'):
        if not isinstance(item, dict):
            continue
        project_id = item.get('project')
        slots.append({
            'project': _project_label(project_id) or _text(project_id),
            'percent': _num(item.get('percentage')),
        })
    return slots


def _advance_slots(bundle):
    """ستون‌های «پیش‌پرداخت مرحله‌ای N»."""
    slots = []
    for item in _attr_list(bundle['contract'], 'advance_payments'):
        if not isinstance(item, dict):
            continue
        slots.append({
            'step': _text(item.get('step')),
            'amount': _num(item.get('amount')),
            'due_date': _text(item.get('due_date')),
            'note': _text(item.get('note')),
        })
    return slots


def _guarantee_item_slots(bundle):
    """ستون‌های «قلم ضمانت قرارداد N»."""
    slots = []
    for item in _attr_list(bundle['contract'], 'guarantee_items'):
        if not isinstance(item, dict):
            continue
        slots.append({
            'type': _text(item.get('type')),
            'amount': _num(item.get('amount')),
            'percent': _num(item.get('percent')),
            'note': _text(item.get('note')),
        })
    return slots


def _dispute_slots(bundle):
    """ستون‌های «اختلاف N»."""
    slots = []
    for dispute in bundle['disputes']:
        slots.append({
            'title': _attr(dispute, 'title'),
            'type': _display(dispute, 'dispute_type'),
            'severity': _display(dispute, 'severity'),
            'status': _display(dispute, 'status'),
            'amount': _num(getattr(dispute, 'claim_amount', None)),
            'opened_date': _fa_date(getattr(dispute, 'opened_date', None)),
            'resolved_date': _fa_date(getattr(dispute, 'resolved_date', None)),
            'resolution': _attr(dispute, 'resolution'),
            'description': _attr(dispute, 'description'),
        })
    return slots


# ----------------------------------------------------------------------------
# مشخصات ستون‌ها: ثابت‌های قرارداد + بلوک‌های آیتم (برای هر آیتم چند ستون)
# ----------------------------------------------------------------------------
# (برچسب بلوک، سازندهٔ برگه‌ها، [(عنوان ستون، کلید مقدار), ...])
BLOCK_SPECS = [
    ('فاکتور', _invoice_slots, [
        ('شماره', 'number'), ('موضوع', 'subject'), ('تاریخ', 'date'), ('سررسید', 'due_date'),
        ('مبلغ', 'amount'), ('مالیات', 'vat'), ('مبلغ کل', 'total'), ('وضعیت', 'status'),
        ('پرداخت‌های مرتبط', 'payments'), ('توضیحات', 'description'),
    ]),
    ('صورت‌وضعیت', _statement_slots, [
        ('شماره', 'number'), ('تاریخ', 'date'), ('تجمعی قبلی', 'cumulative'),
        ('کارکرد دوره', 'work_done'), ('مبلغ', 'amount'), ('مالیات', 'vat'),
        ('جمع اضافات', 'additions_total'), ('جمع کسورات', 'deductions_total'),
        ('مبلغ کل / قابل پرداخت', 'net_amount'), ('وضعیت', 'status'),
        ('پرداخت‌های مرتبط', 'payments'), ('توضیحات', 'description'),
    ]),
    ('اضافهٔ صورت‌وضعیت', lambda b: _statement_extra_slots(b, 'additions', 'اضافه: '), [
        ('صورت‌وضعیت', 'statement'), ('عنوان', 'title'), ('مبلغ', 'amount'),
        ('درصد', 'percent'), ('یادداشت', 'note'),
    ]),
    ('کسور صورت‌وضعیت', lambda b: _statement_extra_slots(b, 'deductions', 'کسور: '), [
        ('صورت‌وضعیت', 'statement'), ('عنوان', 'title'), ('مبلغ', 'amount'),
        ('درصد', 'percent'), ('یادداشت', 'note'),
    ]),
    ('پرداخت', _payment_slots, [
        ('شماره / مرجع', 'reference'), ('تاریخ', 'date'), ('مبلغ', 'amount'),
        ('روش', 'method'), ('مرتبط با', 'linked'), ('یادداشت', 'note'),
    ]),
    ('تضمین', _guarantee_slots, [
        ('شماره', 'number'), ('نوع', 'type'), ('ابزار', 'instrument'),
        ('تاریخ صدور', 'issue_date'), ('تاریخ انقضا', 'expiry_date'), ('مبلغ', 'amount'),
        ('بانک', 'bank'), ('وضعیت', 'status'), ('جزئیات چک/سفته/ضمانت‌نامه', 'details'),
        ('یادداشت', 'note'),
    ]),
    ('الحاقیه', _addendum_slots, [
        ('شماره', 'number'), ('موضوع‌های الحاقیه', 'subjects'), ('موضوع اصلی', 'subject_type'),
        ('شرح موضوع', 'subject'), ('تاریخ', 'date'),
        ('تاریخ پایان جدید', 'new_end_date'), ('تغییر مبلغ', 'amount_change'),
        ('مبلغ جدید قرارداد', 'new_amount'), ('تغییر پیش‌پرداخت', 'advance_change'),
        ('پیش‌پرداخت جدید', 'new_advance_payment'),
        ('درصد حسن انجام کار جدید', 'new_retention_percent'),
        ('مبلغ جدید تضمین', 'new_guarantee_amount'), ('نرخ تسعیر جدید', 'new_exchange_rate'),
        ('تغییر درصد', 'percent_change'), ('تعداد مفاد', 'provisions'),
        ('شرح تغییرات', 'description'),
    ]),
    ('مفاد الحاقیه', _provision_slots, [
        ('الحاقیه / بند', 'addendum'), ('عنوان', 'title'), ('متن بند', 'text'),
    ]),
    ('تخصیص پروژه', _allocation_slots, [
        ('پروژه', 'project'), ('درصد', 'percent'),
    ]),
    ('پیش‌پرداخت مرحله‌ای', _advance_slots, [
        ('مرحله', 'step'), ('مبلغ', 'amount'), ('سررسید', 'due_date'), ('یادداشت', 'note'),
    ]),
    ('قلم ضمانت قرارداد', _guarantee_item_slots, [
        ('نوع', 'type'), ('مبلغ', 'amount'), ('درصد', 'percent'), ('یادداشت', 'note'),
    ]),
    ('اختلاف', _dispute_slots, [
        ('عنوان', 'title'), ('نوع', 'type'), ('شدت', 'severity'), ('وضعیت', 'status'),
        ('مبلغ مطالبه', 'amount'), ('تاریخ باز شدن', 'opened_date'),
        ('تاریخ رفع', 'resolved_date'), ('اقدام / نتیجه', 'resolution'), ('توضیحات', 'description'),
    ]),
]


CONTRACT_HEADERS = [
    'ردیف', 'شماره قرارداد', 'موضوع قرارداد', 'طرف قرارداد', 'نوع شخص طرف', 'نوع قرارداد',
    'طبقه‌بندی', 'وضعیت', 'مبلغ', 'ارز', 'تاریخ شروع', 'تاریخ پایان', 'تاریخ امضا', 'امضاکننده',
    'پیش‌پرداخت', 'درصد حسن انجام کار', 'دوره گارانتی', 'مبلغ تضمین', 'نوع ضمانت',
    # --- مقادیر مؤثر پس از اعمال الحاقیه‌ها ---
    'مبلغ (با الحاقیه)', 'جمع تغییر مبلغ الحاقیه‌ها', 'پیش‌پرداخت (با الحاقیه)',
    'جمع تغییر پیش‌پرداخت الحاقیه‌ها', 'تاریخ پایان (با الحاقیه)',
    'درصد حسن انجام کار (با الحاقیه)', 'مبلغ تضمین (با الحاقیه)',
    'نرخ تسعیر (با الحاقیه)',
    'پروژه‌ها', 'نام پروژه / طرح', 'محل اجرا / تحویل', 'شماره مناقصه / استعلام',
    'تعداد فاکتورها', 'جمع مبلغ فاکتورها', 'تعداد صورت‌وضعیت‌ها', 'جمع مبلغ صورت‌وضعیت‌ها',
    'تعداد پرداخت‌ها', 'جمع مبلغ پرداخت‌ها', 'مانده پرداخت‌نشده', 'درصد پرداخت‌شده',
    'تعداد تضمین‌ها', 'جمع مبالغ تضمین‌ها', 'تعداد الحاقیه‌ها', 'جمع تغییر مبلغ الحاقیه‌ها',
    'تعداد بندهای الحاقیه', 'تعداد اختلافات', 'جمع مبالغ اختلافات',
    'تعداد تخصیص پروژه', 'تعداد پیش‌پرداخت مرحله‌ای', 'تعداد اقلام ضمانت', 'تعداد اسناد پیوست',
    'شرایط و نحوه پرداخت', 'شرایط تحویل', 'وجه التزام / جریمه تأخیر', 'شرایط بیمه',
    'توضیحات قرارداد', 'تلفن طرف', 'ایمیل طرف', 'تاریخ ایجاد', 'آخرین به‌روزرسانی',
]


def _build_headers(payload):
    """سرستون‌ها: ثابت‌های قرارداد + بلوک‌های آیتم بر اساس بیشترین تعداد در داده."""
    headers = list(CONTRACT_HEADERS)
    widths = {}
    for label, builder, columns in BLOCK_SPECS:
        counts = [len(builder(entry['bundle'])) for entry in payload]
        widths[label] = max(counts) if counts else 0
        for slot in range(1, widths[label] + 1):
            for column_label, _key in columns:
                headers.append(f'{label} {slot} — {column_label}')
    return headers, widths


def _contract_row(index, entry, widths):
    """یک ردیف کامل برای یک قرارداد: فیلدهای قرارداد + همهٔ آیتم‌هایش در ستون‌ها."""
    contract = entry['contract']
    bundle = entry['bundle']
    amount = float(getattr(contract, 'amount', None) or 0)
    payment_total = bundle['payment_total']
    balance = amount - payment_total
    paid_percent = round(payment_total / amount * 100, 1) if amount else ''
    addendums = bundle['addendums']
    provisions = sum(len(getattr(a, 'provisions', None) or []) for a in addendums)
    # مقادیر مؤثر پس از اعمال الحاقیه‌ها (بدون بازنویسی دادهٔ قرارداد)
    from contracts.addendum_effects import compute_effects

    effects = compute_effects(contract, addendums)
    eff = effects['effective']
    deltas = effects['deltas']
    row = [
        index, _attr(contract, 'number'), _attr(contract, 'subject'),
        _party_name(contract), _party_type(contract), _type_label(contract),
        _attr(contract, 'category'), _display(contract, 'status'),
        _num(getattr(contract, 'amount', None)), _contract_currency(contract),
        _fa_date(getattr(contract, 'start_date', None)), _fa_date(getattr(contract, 'end_date', None)),
        _fa_date(getattr(contract, 'signing_date', None)),
        _related_name(contract, 'signatory', 'full_name', 'name'),
        _num(getattr(contract, 'advance_payment', None)),
        _num(getattr(contract, 'retention_percent', None)),
        _attr(contract, 'warranty_period'), _num(getattr(contract, 'guarantee_amount', None)),
        _display(contract, 'guarantee_type'),
        # --- مؤثر با الحاقیه ---
        _num(eff.get('amount')), _num(deltas.get('amount')),
        _num(eff.get('advance_payment')), _num(deltas.get('advance_payment')),
        _fa_date(_iso_day(eff.get('end_date'))),
        _num(eff.get('retention_percent')), _num(eff.get('guarantee_amount')),
        _num(eff.get('exchange_rate')),
        _project_names(contract), _attr(contract, 'project_name'),
        _attr(contract, 'project_location'), _attr(contract, 'tender_number'),
        len(bundle['invoices']), _num(bundle['invoice_total']),
        len(bundle['statements']), _num(bundle['statement_total']),
        len(bundle['payments']), _num(payment_total),
        _num(balance), paid_percent,
        len(bundle['guarantees']), _num(bundle['guarantee_total']),
        len(addendums), _num(bundle['addendum_total']),
        provisions, len(bundle['disputes']), _num(bundle['dispute_total']),
        len(_attr_list(contract, 'project_allocations')),
        len(_attr_list(contract, 'advance_payments')),
        len(_attr_list(contract, 'guarantee_items')),
        _related_count(contract, 'documents'),
        _attr(contract, 'payment_terms'), _attr(contract, 'delivery_terms'),
        _attr(contract, 'penalty_terms'), _attr(contract, 'insurance_terms'),
        _attr(contract, 'description'),
        _related_name(contract, 'party', 'phone', 'mobile'),
        _related_name(contract, 'party', 'email'),
        _fa_date(_created_day(contract)), _fa_date(_updated_day(contract)),
    ]
    for label, builder, columns in BLOCK_SPECS:
        slots = builder(bundle)
        for slot in range(widths[label]):
            values = slots[slot] if slot < len(slots) else {}
            for _column_label, key in columns:
                row.append(values.get(key, ''))
    return row


# ----------------------------------------------------------------------------
# کاربرگ یگانه
# ----------------------------------------------------------------------------
MONEY_HINTS = ('مبلغ', 'جمع', 'مالیات', 'پیش‌پرداخت', 'مانده', 'تضمین', 'قابل پرداخت')
TEXT_HINTS = ('تعداد', 'شماره', 'تاریخ', 'درصد', 'وضعیت', 'نوع', 'موضوع', 'عنوان', 'شرح')


def _sheet(workbook, headers):
    """ساخت کاربرگ یگانه (راست‌به‌چپ، سرستون قالب‌بندی‌شده، ردیف سرستون ثابت)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = workbook.active
    ws.title = 'قراردادها'
    ws.sheet_view.rightToLeft = True
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='B45309')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[get_column_letter(index)].width = max(9, min(34, len(str(header)) + 4))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'C2'
    return ws


def _format_columns(ws, headers):
    """قالب اعداد (مبالغ/درصدها) بر اساس عنوان ستون + فعال‌کردن فیلتر سرستون."""
    from openpyxl.utils import get_column_letter

    for index, header in enumerate(headers, start=1):
        header = str(header)
        if 'درصد' in header:
            number_format = '0.0#'
        elif any(hint in header for hint in MONEY_HINTS) and not any(t in header for t in TEXT_HINTS):
            number_format = '#,##0'
        else:
            continue
        for cell in ws[get_column_letter(index)][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format
    last = get_column_letter(max(len(headers), 1))
    ws.auto_filter.ref = f'A1:{last}{max(ws.max_row, 1)}'


def build_contracts_workbook(contracts, **_options):
    """کارپوشهٔ تک‌شیتی با **یک ردیف برای هر قرارداد**.

    تمام آیتم‌های درون هر قرارداد به‌صورت ستون‌های شماره‌دار در همان ردیف می‌آیند؛
    پهنای این ستون‌ها از بیشترین تعداد آیتم در بین قراردادهای خروجی تعیین می‌شود.
    """
    from openpyxl import Workbook

    payload = [{'contract': contract, 'bundle': _bundle(contract)} for contract in contracts]
    headers, widths = _build_headers(payload)

    workbook = Workbook()
    ws = _sheet(workbook, headers)
    for index, entry in enumerate(payload, start=1):
        values = _contract_row(index, entry, widths)
        values = values + [''] * (len(headers) - len(values))
        ws.append(values[:len(headers)])
        ws.cell(row=ws.max_row, column=1).alignment = _center()

    _format_columns(ws, headers)
    return workbook


def _center():
    from openpyxl.styles import Alignment
    return Alignment(horizontal='center', vertical='center')
