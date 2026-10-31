"""سرویس‌های صورت معاملات فصلی (ماده ۱۶۹ مکرر) — نسخهٔ «خرید».

- ``sync_rows``: ذخیرهٔ گروهی و اتمیک ردیف‌های یک جهت معامله (جایگزینی لیست).
- ``validate_report``: اعتبارسنجی قواعدِ خرید از ``seasonal_rules`` (سرور مرجع نهایی).
- ``control_totals``: جمع‌های کنترل یک جهت (مبلغ، مالیات بر ارزش افزوده، جمع).
- ``build_workbook``: خروجی اکسل مدیریتی سه‌شیتی مخصوص خرید.
- ``build_ttms_draft``: بستهٔ «پیش‌نویس آزمایش‌نشده» برای TTMS + راهنمای نگاشت.
"""
from decimal import Decimal, InvalidOperation

from django.db import models, transaction
from django.db.models import Count, Sum

from accounting import seasonal_rules
from accounting.models import SeasonalReport, SeasonalTransaction

DIRECTION_PURCHASE = SeasonalTransaction.Direction.PURCHASE
SUPPORTED_DIRECTIONS = (DIRECTION_PURCHASE,)   # فروش در نسخهٔ بعدی

# ستون‌های خروجی اکسل مدیریتی خرید
PURCHASE_COLUMNS = [
    'ردیف',
    'نوع فروشنده',
    'نام شرکت / عنوان',
    'نام',
    'نام خانوادگی',
    'شناسه ملی / کد ملی',
    'مشمول معاملات فصلی',
    'کد پستی',
    'کشور',
    'استان',
    'شهر',
    'آدرس',
    'شرح کالا / خدمات',
    'نوع کالا یا خدمات',
    'مبلغ خرید (ریال)',
    'مالیات بر ارزش افزوده',
    'جمع کل',
    'نوع معامله',
    'تاریخ فاکتور',
    'شماره فاکتور',
    'وضعیت سامانه مودیان',
]

ROW_FIELDS = (
    'name', 'first_name', 'last_name', 'person_type', 'national_id', 'postal_code',
    'country', 'province', 'city', 'address', 'goods_service_desc', 'amount', 'vat',
    'kind', 'invoice_date', 'invoice_no', 'moadian_status', 'moadian_uid',
    'currency_code', 'fx_rate',
)


def _to_decimal(value):
    try:
        return Decimal(str(value or 0).replace(',', '').strip() or 0)
    except (InvalidOperation, ValueError):
        return Decimal('0')


def _digits_only(value):
    return ''.join(ch for ch in str(value or '') if ch.isdigit())


def _clean_text(value, limit=None):
    text = str(value if value is not None else '').strip()
    return text[:limit] if limit else text


def _parse_date(value):
    """تبدیل تاریخ ISO به ``date``؛ مقدار نامعتبر نادیده گرفته می‌شود."""
    if not value:
        return None
    if hasattr(value, 'year'):
        return value
    try:
        from datetime import date

        return date.fromisoformat(str(value)[:10])
    except (ValueError, TypeError):
        return None


def _ensure_supported(direction):
    if direction not in SUPPORTED_DIRECTIONS:
        raise ValueError('در این نسخه فقط معاملات خرید پشتیبانی می‌شود.')
    return direction


def sync_rows(report, rows, direction=DIRECTION_PURCHASE):
    """جایگزینی کامل ردیف‌های یک جهت معامله در گزارش (تراکنشی).

    ردیف‌های جهت دیگر (مثلاً فروش در آینده) دست‌نخورده می‌مانند و شماره‌گذاری
    ردیف‌ها به‌ازای هر جهت مستقل است.
    """
    direction = _ensure_supported(direction)
    if report.is_locked:
        raise ValueError('این فصل ارسال شده است و قابل ویرایش نیست.')

    cleaned = []
    for index, raw in enumerate(rows or [], start=1):
        if not isinstance(raw, dict):
            continue
        item = {
            'id': raw.get('id') or None,
            'line_no': index,
            'person_type': _clean_text(raw.get('person_type'), 10),
            'first_name': _clean_text(raw.get('first_name'), 150),
            'last_name': _clean_text(raw.get('last_name'), 150),
            'name': _clean_text(raw.get('name'), 250),
            'national_id': _digits_only(raw.get('national_id'))[:20],
            'postal_code': _digits_only(raw.get('postal_code'))[:10],
            'country': _clean_text(raw.get('country'), 60),
            'province': _clean_text(raw.get('province'), 60),
            'city': _clean_text(raw.get('city'), 100),
            'address': _clean_text(raw.get('address')),
            'goods_service_desc': _clean_text(raw.get('goods_service_desc'), 300),
            'amount': _to_decimal(raw.get('amount')),
            'vat': _to_decimal(raw.get('vat')),
            'kind': int(raw.get('kind') or 1) if str(raw.get('kind') or '1').isdigit() else 1,
            'invoice_date': _parse_date(raw.get('invoice_date')),
            'invoice_no': _clean_text(raw.get('invoice_no'), 50),
            'moadian_status': 'registered' if raw.get('moadian_status') == 'registered' else 'pending',
            'moadian_uid': _clean_text(raw.get('moadian_uid'), 120),
            'currency_code': _clean_text(raw.get('currency_code'), 10),
            'fx_rate': _to_decimal(raw.get('fx_rate')) or Decimal('1'),
        }
        if item['kind'] < 1 or item['kind'] > 5:
            item['kind'] = 1
        # شخص حقیقی: نام و نام خانوادگی جدا ذخیره می‌شود و «نام کامل» برای گزارش ساخته می‌شود
        if item['person_type'] == SeasonalTransaction.PersonType.NATURAL:
            composed = ' '.join(part for part in [item['first_name'], item['last_name']] if part).strip()
            if composed:
                item['name'] = composed[:250]
        # کشور پیش‌فرض «ایران» برای همهٔ انواع به‌جز خارجی (که کاربر باید وارد کند)
        if item['person_type'] != SeasonalTransaction.PersonType.FOREIGN and not item['country']:
            item['country'] = seasonal_rules.DEFAULT_COUNTRY
        if not any([item['name'], item['first_name'], item['last_name'], item['invoice_no'],
                    item['amount'], item['vat'], item['goods_service_desc']]):
            continue
        cleaned.append(item)

    scoped = report.transactions.filter(direction=direction)
    with transaction.atomic():
        existing = {str(row.id): row for row in scoped}
        # شماره‌های موقت تا محدودیت یکتایی (گزارش، جهت، ردیف) در بازچینش نقض نشود
        scoped.update(line_no=models.F('line_no') + 100000)
        keep_ids = []
        for item in cleaned:
            row_id = item.pop('id', None)
            row = existing.get(str(row_id)) if row_id else None
            if row is None:
                row = SeasonalTransaction(
                    report=report, company_id=report.company_id, direction=direction,
                )
            for field in ROW_FIELDS:
                setattr(row, field, item[field])
            # دستور کارفرما: همهٔ ردیف‌ها مشمول و نوع کالا/خدمات «سایر»
            row.is_included = True
            row.goods_service_type = SeasonalTransaction.GoodsServiceType.OTHER
            row.direction = direction
            row.line_no = item['line_no']
            row.save()
            keep_ids.append(row.id)
        scoped.exclude(id__in=keep_ids).delete()

    return scoped.count()

def _scoped_rows(report, direction=DIRECTION_PURCHASE):
    """ردیف‌های یک جهت معامله در گزارش."""
    return report.transactions.filter(direction=direction)


def control_totals(report, direction=DIRECTION_PURCHASE):
    """جمع‌های کنترل یک جهت معامله: مبلغ، مالیات بر ارزش افزوده و جمع کل."""
    rows = _scoped_rows(report, direction)
    agg = rows.aggregate(
        count=Count('id'),
        amount=Sum('amount'),
        vat=Sum('vat'),
    )
    by_kind = [
        {
            'kind': item['kind'],
            'label': str(SeasonalTransaction.Kind(item['kind']).label),
            'count': item['count'],
            'amount': float(item['amount'] or 0),
            'vat': float(item['vat'] or 0),
            'total': float((item['amount'] or 0) + (item['vat'] or 0)),
        }
        for item in (
            rows.order_by()          # پاک‌کردن ordering پیش‌فرض مدل تا گروه‌بندی درست باشد
            .values('kind')
            .annotate(count=Count('id'), amount=Sum('amount'), vat=Sum('vat'))
            .order_by('kind')
        )
    ]
    registered = rows.filter(moadian_status=SeasonalTransaction.MoadianStatus.REGISTERED).count()
    amount = float(agg['amount'] or 0)
    vat = float(agg['vat'] or 0)
    by_person_type = [
        {
            'person_type': item['person_type'] or 'unknown',
            'label': dict(seasonal_rules.PERSON_TYPES).get(item['person_type'], 'نامشخص'),
            'count': item['count'],
            'amount': float(item['amount'] or 0),
        }
        for item in (
            rows.order_by()
            .values('person_type')
            .annotate(count=Count('id'), amount=Sum('amount'))
            .order_by('person_type')
        )
    ]
    return {
        'direction': direction,
        'count': agg['count'] or 0,
        'amount': amount,
        'vat': vat,
        'total': amount + vat,
        'registered': registered,
        'pending': (agg['count'] or 0) - registered,
        'by_kind': by_kind,
        'by_person_type': by_person_type,
    }


def validate_report(report, direction=DIRECTION_PURCHASE):
    """اعتبارسنجی ردیف‌های خرید بر پایهٔ قواعد اعلانی ``seasonal_rules``.

    خروجی: ``{'issues': [...], 'summary': {...}, 'ruleset': {...}}``.
    هر issue دارای ``code``/``source`` است تا «الزام قانونی» از «الزام نسخهٔ
    TTMS تأییدنشده» و «سیاست داخلی» قابل تفکیک باشد. اعتبارسنجی سرور مرجع نهایی است.
    """
    rows = list(_scoped_rows(report, direction).order_by('line_no', 'id'))
    period_start, period_end = report.period_range
    issues = []
    seen = {}

    def add(row, field, level, message, code='', source=None):
        issues.append({
            'row_id': row.id,
            'row_no': row.line_no,
            'row_name': row.name,
            'field': field,
            'field_label': seasonal_rules.FIELD_LABELS.get(field, field),
            'level': level,
            'message': message,
            'code': code,
            'source': source or seasonal_rules.SOURCE_INTERNAL,
        })

    for row in rows:
        person_type = row.person_type or ''
        # ---- قواعد اعلانی: required / recommended / digits / forbidden ----
        for rule in seasonal_rules.rules_for(person_type or None):
            if person_type and person_type not in rule['person_types']:
                continue
            value = getattr(row, rule['field'], '')
            text = str(value if value is not None else '').strip()
            kind = rule['kind']
            if kind in ('required', 'recommended'):
                if not text:
                    add(row, rule['field'], rule['level'], rule['message'], rule['code'], rule['source'])
            elif kind == 'digits':
                expected = int(rule.get('length') or 0)
                if text and (not text.isdigit() or (expected and len(text) != expected)):
                    add(row, rule['field'], rule['level'], rule['message'], rule['code'], rule['source'])
            elif kind == 'forbidden':
                if text:
                    add(row, rule['field'], rule['level'], rule['message'], rule['code'], rule['source'])

        # ---- قواعد محاسباتی و میان‌فیلدی خرید (فقط مالیات بر ارزش افزوده) ----
        if not row.amount and not row.vat:
            add(row, 'amount', 'error', 'مبلغ خرید باید بزرگ‌تر از صفر باشد.',
                'amount_positive', seasonal_rules.SOURCE_LEGAL)
        if row.vat and row.amount and row.vat > row.amount:
            add(row, 'vat', 'error', 'مالیات بر ارزش افزوده نمی‌تواند از مبلغ خرید بیشتر باشد.',
                'vat_le_amount', seasonal_rules.SOURCE_LEGAL)
        if row.invoice_date and not (period_start <= row.invoice_date <= period_end):
            add(row, 'invoice_date', 'error', 'تاریخ فاکتور بیرون از بازهٔ این فصل است.',
                'invoice_date_in_period', seasonal_rules.SOURCE_LEGAL)
        # ردیف «تجمیعی» فقط مبلغ دارد؛ هشدارهای فاکتور/سامانه برای آن معنا ندارد
        if row.person_type != SeasonalTransaction.PersonType.AGGREGATE:
            if row.moadian_status != SeasonalTransaction.MoadianStatus.REGISTERED:
                add(row, 'moadian_status', 'warning', 'این فاکتور در سامانه مودیان درج نشده است.',
                    'moadian_not_registered', seasonal_rules.SOURCE_TTMS)
            elif not row.moadian_uid:
                add(row, 'moadian_uid', 'warning',
                    'برای فاکتور درج‌شده در سامانه مودیان، شناسهٔ یکتا بهتر است ثبت شود.',
                    'moadian_uid_recommended', seasonal_rules.SOURCE_TTMS)

        if row.invoice_no:
            key = (row.invoice_no.strip(), row.national_id.strip() or row.name.strip())
            if key in seen:
                add(row, 'invoice_no', 'error',
                    f'شماره فاکتور تکراری با ردیف {seen[key]} ثبت شده است.',
                    'duplicate_invoice', seasonal_rules.SOURCE_INTERNAL)
            else:
                seen[key] = row.line_no

    errors = sum(1 for issue in issues if issue['level'] == 'error')
    warnings = sum(1 for issue in issues if issue['level'] == 'warning')
    return {
        'issues': issues,
        'direction': direction,
        'ruleset': {
            'version': seasonal_rules.RULESET_VERSION,
            'status': seasonal_rules.RULESET_STATUS,
        },
        'summary': {
            'rows': len(rows),
            'errors': errors,
            'warnings': warnings,
            'period_start': period_start.isoformat(),
            'period_end': period_end.isoformat(),
            'totals': control_totals(report, direction),
        },
    }


# ---------------------------------------------------------------------------
# خروجی اکسل رسمی (سه شیت: داده / خلاصه / خطاها)
# ---------------------------------------------------------------------------

def _jalali_text(value):
    """تاریخ میلادی → رشتهٔ شمسی (مثل ۱۴۰۵/۰۱/۲۱)."""
    if not value:
        return ''
    import jdatetime

    jalali = jdatetime.date.fromgregorian(date=value)
    return f'{jalali.year:04d}/{jalali.month:02d}/{jalali.day:02d}'


def _company_identity(report):
    """نام و شناسهٔ قانونی شرکت برای سرصفحهٔ گزارش."""
    name = getattr(report.company, 'name', '') or ''
    economic_code = ''
    national_id = ''
    try:
        from settings_app.models import CompanyProfile

        profile = CompanyProfile.objects.filter(company_id=report.company_id).first()
        if profile:
            name = profile.legal_name or name
            economic_code = profile.economic_code or ''
            national_id = profile.national_id or ''
    except Exception:                      # تعریف‌نشده/خطای دسترسی — سرصفحه ساده می‌ماند
        pass
    return name, economic_code, national_id


def _row_values(row):
    """مقادیر یک ردیف خرید به ترتیب ستون‌های ``PURCHASE_COLUMNS``."""
    person_label = dict(seasonal_rules.PERSON_TYPES).get(row.person_type, '')
    return [
        row.line_no,
        person_label,
        row.name,
        row.first_name,
        row.last_name,
        row.national_id,
        'بله' if row.is_included else 'خیر',
        row.postal_code,
        row.country,
        row.province,
        row.city,
        row.address,
        row.goods_service_desc,
        str(row.get_goods_service_type_display()),
        int(row.amount or 0),
        int(row.vat or 0),
        int(row.total or 0),
        int(row.kind),
        _jalali_text(row.invoice_date),
        row.invoice_no,
        'درج‌شده در سامانه' if row.moadian_status == 'registered' else 'درج‌نشده',
    ]


def _build_data_sheet(workbook, report, rows, totals):
    """شیت «داده» خرید: سرصفحه + ردیف‌ها + سطر جمع (مبلغ/ارزش افزوده/جمع)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color='FFFFFF', size=10)
    head_fill = PatternFill('solid', fgColor='B45309')
    total_fill = PatternFill('solid', fgColor='FEF3C7')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    thin = Side(style='thin', color='D8DEE9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    columns = len(PURCHASE_COLUMNS)

    ws = workbook.active
    ws.title = 'داده'
    ws.sheet_view.rightToLeft = True

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    title = ws.cell(row=1, column=1, value='صورت معاملات فصلی خرید — ماده ۱۶۹ مکرر قانون مالیات‌های مستقیم')
    title.font = Font(bold=True, size=13)
    title.alignment = center

    company_name, company_economic, company_national = _company_identity(report)
    period_start, period_end = report.period_range
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    meta = ws.cell(row=2, column=1, value=(
        f'{company_name}  |  شناسه ملی: {company_national or "—"}  |  شماره اقتصادی: {company_economic or "—"}'
        f'  |  سال مالی: {report.fiscal_year.name}  |  {report.get_quarter_display()}'
        f'  |  بازه: {_jalali_text(period_start)} تا {_jalali_text(period_end)}'
        f'  |  قواعد: {seasonal_rules.RULESET_VERSION} ({seasonal_rules.RULESET_STATUS})'
    ))
    meta.alignment = center

    for index, header in enumerate(PURCHASE_COLUMNS, start=1):
        cell = ws.cell(row=3, column=index, value=header)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(index)].width = max(10, min(30, len(header) + 6))
    ws.row_dimensions[3].height = 34
    ws.freeze_panes = 'A4'

    for offset, row in enumerate(rows):
        excel_row = 4 + offset
        for column, value in enumerate(_row_values(row), start=1):
            cell = ws.cell(row=excel_row, column=column, value=value)
            cell.border = border
            if column in (6, 8, 19, 20):        # شناسه/کد پستی/تاریخ/شماره فاکتور به‌صورت متن
                cell.number_format = '@'
                cell.alignment = center
            elif column in (1, 2, 7, 14, 18, 21):
                cell.alignment = center
            elif column in (15, 16, 17):
                cell.number_format = '#,##0'
                cell.alignment = left

    total_row = 4 + len(rows)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=14)
    label = ws.cell(row=total_row, column=1, value='جمع کل خرید')
    label.font = Font(bold=True)
    label.alignment = center
    for column, value in (
        (15, totals['amount']), (16, totals['vat']), (17, totals['total']),
    ):
        cell = ws.cell(row=total_row, column=column, value=value)
        cell.font = Font(bold=True)
        cell.number_format = '#,##0'
        cell.alignment = left
    for column in range(1, columns + 1):
        cell = ws.cell(row=total_row, column=column)
        cell.fill = total_fill
        cell.border = border

    if rows:
        last_col = get_column_letter(columns)
        ws.auto_filter.ref = f'A3:{last_col}{3 + len(rows)}'
    return ws

def _build_summary_sheet(workbook, report, totals, validation):
    """شیت «خلاصه» خرید: جمع‌ها، به تفکیک نوع معامله و نوع فروشنده + وضعیت قواعد."""
    from openpyxl.styles import Font, PatternFill

    head_font = Font(bold=True, color='FFFFFF', size=10)
    head_fill = PatternFill('solid', fgColor='B45309')

    period_start, period_end = report.period_range
    ws = workbook.create_sheet('خلاصه')
    ws.sheet_view.rightToLeft = True

    ws.append(['صورت معاملات فصلی خرید — خلاصهٔ کنترلی'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    for label, value in (
        ('سال مالی', report.fiscal_year.name),
        ('فصل', report.get_quarter_display()),
        ('بازهٔ فصل', f'{_jalali_text(period_start)} تا {_jalali_text(period_end)}'),
        ('وضعیت گزارش', report.get_status_display()),
        ('نسخهٔ قواعد اعتبارسنجی', seasonal_rules.RULESET_VERSION),
        ('وضعیت تأیید قواعد/قالب', seasonal_rules.RULESET_STATUS),
        ('', ''),
        ('تعداد ردیف خرید', totals['count']),
        ('جمع مبلغ خرید (ریال)', totals['amount']),
        ('جمع مالیات بر ارزش افزوده', totals['vat']),
        ('جمع کل', totals['total']),
        ('درج‌شده در سامانه مودیان', totals['registered']),
        ('درج‌نشده در سامانه', totals['pending']),
        ('خطاهای اعتبارسنجی', validation['summary']['errors']),
        ('هشدارهای اعتبارسنجی', validation['summary']['warnings']),
    ):
        ws.append([label, value])
    ws.append([])
    ws.append(['نوع معامله', 'تعداد', 'مبلغ', 'مالیات بر ارزش افزوده', 'جمع کل'])
    for cell in ws[ws.max_row]:
        cell.font = head_font
        cell.fill = head_fill
    for item in totals['by_kind']:
        ws.append([
            item['label'], item['count'], item['amount'],
            item['vat'], item['total'],
        ])
    ws.append([])
    ws.append(['نوع فروشنده', 'تعداد', 'مبلغ'])
    for cell in ws[ws.max_row]:
        cell.font = head_font
        cell.fill = head_fill
    for item in totals.get('by_person_type', []):
        ws.append([item['label'], item['count'], item['amount']])
    for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row_cells:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
    ws.column_dimensions['A'].width = 32
    for letter in ('B', 'C', 'D', 'E', 'F'):
        ws.column_dimensions[letter].width = 18
    return ws


def _build_issues_sheet(workbook, validation):
    """شیت «خطاها»: فهرست خطا/هشدارها با کد قاعده و منبع (قابل ردیابی)."""
    from openpyxl.styles import Alignment, Font, PatternFill

    head_font = Font(bold=True, color='FFFFFF', size=10)
    head_fill = PatternFill('solid', fgColor='B45309')

    ws = workbook.create_sheet('خطاها')
    ws.sheet_view.rightToLeft = True
    ws.append(['ردیف', 'فروشنده', 'فیلد', 'سطح', 'پیام', 'کد قاعده', 'منبع قاعده'])
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for issue in validation['issues']:
        ws.append([
            issue['row_no'], issue['row_name'], issue.get('field_label') or issue['field'],
            'خطا' if issue['level'] == 'error' else 'هشدار', issue['message'],
            issue.get('code', ''), issue.get('source', ''),
        ])
    if not validation['issues']:
        ws.append(['', 'موردی برای اصلاح نیست.', '', '', '', '', ''])
    for column, width in zip('ABCDEFG', (8, 30, 18, 10, 55, 26, 50)):
        ws.column_dimensions[column].width = width
    return ws


def build_workbook(report, direction=DIRECTION_PURCHASE):
    """کارپوشهٔ خروجی مدیریتی خرید: داده + خلاصه + خطاها (فقط ردیف‌های خرید)."""
    from openpyxl import Workbook

    direction = _ensure_supported(direction)
    validation = validate_report(report, direction)
    rows = list(_scoped_rows(report, direction).order_by('line_no', 'id'))
    totals = validation['summary']['totals']

    workbook = Workbook()
    _build_data_sheet(workbook, report, rows, totals)
    _build_summary_sheet(workbook, report, totals, validation)
    _build_issues_sheet(workbook, validation)
    return workbook



# ---------------------------------------------------------------------------
# خروجی TTMS (پیش‌نویس آزمایش‌نشده)
# ---------------------------------------------------------------------------

def _draft_value(row, internal):
    """مقدار یک ردیف برای پیش‌نویس TTMS بر اساس فیلد داخلی."""
    if internal == 'person_type':
        return dict(seasonal_rules.PERSON_TYPES).get(row.person_type, '')
    if internal == 'is_included':
        return 'بله' if row.is_included else 'خیر'
    if internal == 'goods_service_type':
        return str(row.get_goods_service_type_display())
    if internal == 'moadian_status':
        return 'درج‌شده در سامانه' if row.moadian_status == 'registered' else 'درج‌نشده'
    if internal == 'invoice_date':
        return _jalali_text(row.invoice_date)
    if internal in ('amount', 'vat', 'total', 'kind', 'line_no'):
        value = getattr(row, internal, None)
        return int(value or 0)
    return str(getattr(row, internal, '') or '')


def build_ttms_draft(report, direction=DIRECTION_PURCHASE):
    """ساخت بستهٔ «پیش‌نویس TTMS» — **آزمایش‌نشده** (قالب رسمی در دسترس نیست).

    خروجی: ``(zip_bytes, filename, warnings)``. بسته شامل:
    - ``purchase-rows.csv``: فقط ردیف‌های خرید، با ستون‌های نگاشتِ پیش‌نویس
    - ``README.txt``: هشدار عدم تأیید، نسخهٔ قواعد، جدول نگاشت و اقلام لازم برای تکمیل
    """
    import csv
    import io
    import zipfile

    direction = _ensure_supported(direction)
    rows = list(_scoped_rows(report, direction).order_by('line_no', 'id'))
    mapping = seasonal_rules.TTMS_DRAFT_MAPPING
    totals = control_totals(report, direction)
    period_start, period_end = report.period_range
    company_name, company_economic, company_national = _company_identity(report)

    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator='\r\n')
    writer.writerow([seasonal_rules.FIELD_LABELS.get(item['internal'], item['internal']) for item in mapping])
    for row in rows:
        writer.writerow([_draft_value(row, item['internal']) for item in mapping])
    csv_text = '\ufeff' + buffer.getvalue()

    mapping_lines = [
        f"- {seasonal_rules.FIELD_LABELS.get(item['internal'], item['internal'])}"
        f" | فیلد داخلی: {item['internal']}"
        f" | نوع: {item['type']}"
        f" | طول: {item.get('length') or '—'}"
        f" | الزام: {item.get('requirement', '—')}"
        f" | نام فیلد در قالب رسمی: {item.get('target_field') or 'نامشخص (نیازمند قالب رسمی)'}"
        for item in mapping
    ]
    readme = '\r\n'.join([
        'بستهٔ پیش‌نویس خروجی TTMS — وضعیت: آزمایش‌نشده (UNVERIFIED)',
        '=' * 70,
        '',
        'هشدار مهم:',
        'قالب رسمی/Access نسخهٔ TTMS هدف و راهنمای آن در محیط پروژه موجود نبود؛',
        'این فایل صرفاً یک «پیش‌نویس ساختاری» است و نباید به‌عنوان خروجی تأییدشده',
        'یا سازگار قطعی با نرم‌افزار TTMS تلقی شود. پیش از استفادهٔ رسمی، فایل باید',
        'با نسخهٔ مشخص TTMS آزمایش و نگاشت ستون‌ها نهایی شود.',
        '',
        f'نسخهٔ قواعد اعتبارسنجی: {seasonal_rules.RULESET_VERSION} ({seasonal_rules.RULESET_STATUS})',
        '',
        'سرصفحهٔ گزارش:',
        f'  شرکت: {company_name}',
        f'  شناسه ملی شرکت: {company_national or "—"}',
        f'  شماره اقتصادی شرکت: {company_economic or "—"}',
        f'  سال مالی: {report.fiscal_year.name} | {report.get_quarter_display()}',
        f'  بازهٔ فصل: {_jalali_text(period_start)} تا {_jalali_text(period_end)}',
        f'  جهت معامله: خرید',
        '',
        'جمع‌های کنترل (فقط خرید):',
        f'  تعداد ردیف: {totals["count"]}',
        f'  جمع مبلغ: {totals["amount"]}',
        f'  جمع مالیات بر ارزش افزوده: {totals["vat"]}',
        f'  جمع کل: {totals["total"]}',
        '',
        'نگاشت ستون‌های پیش‌نویس:',
        *mapping_lines,
        '',
        'قواعد ثابت این نسخه (دستور کارفرما):',
        '  - همهٔ ردیف‌ها «مشمول معاملات فصلی» = بله',
        '  - «نوع کالا یا خدمات» = سایر',
        '  - فروشندهٔ حقیقی: نام و نام خانوادگی جدا + کد ملی ۱۰ رقمی',
        '  - فروشندهٔ حقوقی: شناسه ملی ۱۱ رقمی (کد اقتصادی استفاده نمی‌شود)',
        '',
        'برای تکمیل و تأیید این خروجی، موارد زیر لازم است:',
        '  1) نسخهٔ دقیق نرم‌افزار آفلاین TTMS مورد استفاده در مؤسسه برای سال ۱۴۰۵',
        '  2) فایل خام Access (یا پایگاه‌دادهٔ رسمی) همان نسخه و راهنمای همراه آن',
        '  3) نام/نوع/طول/مقدار مجاز هر ستون در بخش «خرید» و تفاوت الزامات فروشندهٔ',
        '     حقیقی/حقوقی/خارجی در همان نسخه',
        '  4) نمونهٔ فایل ورودی موفق (برای آزمون رفت‌وبرگشت)',
        '',
        'پس از دریافت موارد بالا، تنها نگاشت seasonal_rules.TTMS_DRAFT_MAPPING و',
        'سطح/شدت قواعد همان فایل به‌روزرسانی می‌شود؛ ساختار فرم و API تغییر نمی‌کند.',
    ])

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('purchase-rows.csv', csv_text)
        archive.writestr('README.txt', readme)

    warnings = [
        'خروجی TTMS «آزمایش‌نشده» است: قالب رسمی نسخهٔ هدف در دسترس نبود.',
        'نام ستون‌های رسمی TTMS نگاشت نشده و در README به‌صورت «نامشخص» علامت خورده است.',
        'پیش از ارسال رسمی، فایل باید با نسخهٔ مشخص TTMS آزمایش شود.',
    ]
    year_key = getattr(report, 'fiscal_year_id', None) or getattr(report.fiscal_year, 'pk', 'na')
    filename = f'ttms-draft-purchase-{year_key}-q{report.quarter}.zip'
    return zip_buffer.getvalue(), filename, warnings

