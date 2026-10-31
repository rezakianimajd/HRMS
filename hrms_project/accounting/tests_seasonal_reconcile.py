"""تست‌های گزارش‌گیری و مغایرت‌گیری «اکسل دفاتر» برای معاملات فصلی (خرید).

اجرای واقعی روی سرور: ``manage.py test accounting.tests_seasonal_reconcile``.
"""
import io
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile

from accounting import seasonal_reconcile as rec
from accounting.seasonal import sync_rows
from accounting.tests_seasonal import SeasonalBaseTest


class ReconcileHelperTests(SeasonalBaseTest):
    """نرمال‌سازی متن/عدد/تاریخ و شباهت نام."""

    def test_digits_ids_and_invoices(self):
        self.assertEqual(rec.normalize_id('۱۰۱۰۱۰۱۰۱۰۱'), '10101010101')
        self.assertEqual(rec.normalize_id('101-010-10101'), '10101010101')
        self.assertEqual(rec.normalize_invoice('۰۰۰۱۰۰۱'), '1001')
        self.assertEqual(rec.normalize_invoice('0'), '')
        self.assertEqual(rec.normalize_invoice(''), '')

    def test_amounts_and_dates(self):
        self.assertEqual(rec.parse_amount('۱٬۲۵۰٬۰۰۰ ریال'), Decimal('1250000'))
        self.assertEqual(rec.parse_amount('(1,200)'), Decimal('-1200'))
        self.assertEqual(rec.parse_amount(None), Decimal('0'))
        self.assertEqual(rec.parse_date('1405/01/21').isoformat(), '2026-04-10')
        self.assertEqual(rec.parse_date('2026-04-10').isoformat(), '2026-04-10')
        self.assertIsNone(rec.parse_date('بدون تاریخ'))

    def test_name_similarity_ignores_suffixes_and_order(self):
        self.assertGreater(rec.name_similarity('شرکت آلفا تجارت', 'آلفا تجارت'), 0.9)
        self.assertGreater(rec.name_similarity('شرکت الفا تجارت', 'شرکت آلفا تجارت'), 0.8)
        self.assertLess(rec.name_similarity('فروشگاه گاما', 'شرکت بتا'), 0.4)


class LedgerReaderTests(SeasonalBaseTest):
    """خواندن فایل دفاتر (CSV و اکسل) و تشخیص ستون‌ها."""

    def test_csv_detects_columns_and_rows(self):
        content = (
            '\ufeffشماره فاکتور,تاریخ فاکتور,نام فروشنده,شناسه ملی,مبلغ,مالیات\n'
            '1001,1405/01/21,شرکت آلفا تجارت,10101010101,1000000,100000\n'
            ',,,,\n'
        ).encode('utf-8')
        parsed = rec.read_ledger_rows(SimpleUploadedFile('ledger.csv', content))
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.mapping['invoice_no'], 0)
        self.assertEqual(parsed.rows[0]['invoice_no_key'], '1001')
        self.assertEqual(parsed.rows[0]['amount'], Decimal('1000000'))
        self.assertEqual(parsed.rows[0]['invoice_date'].isoformat(), '2026-04-10')

    def test_xlsx_header_is_detected_in_third_row(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['دفتر خرید'])
        sheet.append([])
        sheet.append(['تاریخ سند', 'شماره سند', 'طرف حساب', 'مبلغ'])
        sheet.append(['1405/01/21', '1001', 'شرکت آلفا تجارت', 1000000])
        stream = io.BytesIO()
        workbook.save(stream)
        stream.seek(0)

        parsed = rec.read_ledger_rows(SimpleUploadedFile('ledger.xlsx', stream.getvalue()))
        self.assertEqual(parsed.header_row, 3)
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.rows[0]['party_name'], 'شرکت آلفا تجارت')

    def test_unsupported_format_is_rejected(self):
        with self.assertRaises(ValueError):
            rec.read_ledger_rows(SimpleUploadedFile('ledger.xls', b'binary'))


class ReconcileEngineTests(SeasonalBaseTest):
    """موتور تطبیق: تطبیق کامل، مغایرت مبلغ، فقط یک طرف و پیشنهاد افزودن ردیف."""

    def setUp(self):
        super().setUp()
        self.report = self.make_report()
        sync_rows(self.report, [
            self.row(invoice_no='1001', amount=1000000, vat=100000),
            self.row(invoice_no='1002', name='شرکت بتا صنعت', national_id='20202020202',
                     amount=500000, vat=50000),
            self.row(invoice_no='1003', name='شرکت گاما', national_id='30303030303',
                     amount=300000, vat=0),
        ])

    def ledger(self):
        def line(index, file_row, invoice, name, national_id, amount, vat):
            return {
                'index': index, 'file_row': file_row, 'invoice_no': invoice,
                'invoice_no_key': invoice, 'invoice_date': None, 'party_name': name,
                'name_key': rec.normalize_name(name), 'national_id': national_id,
                'amount': Decimal(amount), 'vat': Decimal(vat),
                'total': Decimal(amount) + Decimal(vat),
            }

        return [
            line(1, 2, '1001', 'شرکت آلفا تجارت', '10101010101', '1000000', '100000'),
            line(2, 3, '1002', 'شرکت بتا صنعت', '20202020202', '510000', '50000'),
            line(3, 4, '1004', 'شرکت دلتا', '40404040404', '700000', '70000'),
        ]

    def test_matching_classifies_rows(self):
        result = rec.reconcile(self.report, self.ledger(), rec.MatchOptions(), 'purchase')
        summary = result['summary']
        self.assertEqual(summary['seasonal']['count'], 3)
        self.assertEqual(summary['matched']['count'], 2)
        self.assertEqual(summary['counts']['amount_mismatch'], 1)
        self.assertEqual(summary['counts']['missing_in_seasonal'], 1)
        self.assertEqual(summary['counts']['missing_in_ledger'], 1)
        self.assertEqual(summary['amount_delta'], float(1800000 - 2210000))
        self.assertEqual(result['suggested_rows'][0]['invoiceNo'], '1004')

    def test_amount_mismatch_has_evidence(self):
        result = rec.reconcile(self.report, self.ledger(), rec.MatchOptions(), 'purchase')
        mismatch = [item for item in result['items'] if item['kind'] == 'amount_mismatch'][0]
        self.assertEqual(mismatch['seasonal']['invoice_no'], '1002')
        self.assertTrue(mismatch['evidence']['invoice_exact'])
        self.assertEqual(mismatch['evidence']['amount_state'], 'different')
        self.assertTrue(any(diff['field'] == 'amount' for diff in mismatch['differences']))

    def test_tolerance_option_softens_minor_differences(self):
        options = rec.MatchOptions.from_payload({'amount_tolerance': 15000})
        result = rec.reconcile(self.report, self.ledger(), options, 'purchase')
        self.assertEqual(result['summary']['counts']['amount_mismatch'], 0)
        self.assertEqual(result['summary']['differences']['amount_minor'], 1)

    def test_name_only_match_without_invoice_number(self):
        rows = [{'index': 1, 'file_row': 2, 'invoice_no': '', 'invoice_no_key': '',
                 'invoice_date': None, 'party_name': 'شرکت گاما',
                 'name_key': rec.normalize_name('شرکت گاما'), 'national_id': '',
                 'amount': Decimal('300000'), 'vat': Decimal('0'), 'total': Decimal('300000')}]
        result = rec.reconcile(self.report, rows, rec.MatchOptions(), 'purchase')
        matched = [item for item in result['items'] if item['kind'] in ('matched', 'difference')]
        self.assertTrue(matched)
        self.assertEqual(matched[0]['seasonal']['invoice_no'], '1003')
        self.assertFalse(matched[0]['evidence']['invoice_exact'])

    def test_reconciliation_workbook_layout(self):
        result = rec.reconcile(self.report, self.ledger(), rec.MatchOptions(), 'purchase')
        workbook = rec.build_reconciliation_workbook(result)
        self.assertEqual(workbook.sheetnames, [
            'خلاصه', 'مغایرت‌ها', 'فقط در فصلی', 'فقط در دفاتر', 'تطبیق‌شده', 'نگاشت ستون‌ها',
        ])
        self.assertEqual(workbook['خلاصه'].cell(row=1, column=1).value,
                         'گزارش مغایرت‌گیری صورت معاملات فصلی خرید')
        self.assertEqual(workbook['فقط در دفاتر'].cell(row=1, column=2).value, 'نوع مغایرت')
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.assertGreater(buffer.tell(), 3000)


class AnalyticsTests(SeasonalBaseTest):
    """شاخص‌های تحلیلی فصل."""

    def test_analytics_filters_by_kind_and_status(self):
        report = self.make_report()
        sync_rows(report, [
            self.row(invoice_no='K1', amount=1000000, vat=100000, kind=1),
            self.row(invoice_no='K2', name='شرکت بتا', national_id='20202020202',
                     amount=500000, vat=50000, kind=2, moadian_status='pending'),
            self.row(invoice_no='K3', name='شرکت گاما', national_id='30303030303',
                     amount=300000, vat=0, kind=2),
        ])

        only_kind2 = rec.build_analytics(report, 'purchase', {'kind': 2})
        self.assertEqual(only_kind2['totals']['count'], 2)
        self.assertEqual(only_kind2['totals']['amount'], 800000.0)
        self.assertTrue(only_kind2['active_filters'])
        self.assertEqual(only_kind2['unfiltered']['count'], 3)
        self.assertEqual(only_kind2['filters']['kind'], 2)

        pending = rec.build_analytics(report, 'purchase', {'moadian_status': 'pending'})
        self.assertEqual(pending['totals']['count'], 1)
        self.assertEqual(pending['totals']['amount'], 500000.0)

        combined = rec.build_analytics(report, 'purchase', {'kind': '2', 'moadian_status': 'pending'})
        self.assertEqual(combined['totals']['count'], 1)

        self.assertEqual(len(rec.filtered_rows(report, 'purchase', {'kind': 2})), 2)
        self.assertIsNone(rec.filters_from_query({'kind': '۹'})['kind'])

    def test_filter_options_and_analytics_workbook(self):
        report = self.make_report()
        sync_rows(report, [
            self.row(invoice_no='W1', amount=1000000, vat=100000, kind=1),
            self.row(invoice_no='W2', name='شرکت بتا', national_id='20202020202',
                     amount=500000, vat=50000, kind=2, moadian_status='pending'),
        ])
        analytics = rec.build_analytics(report, 'purchase', {'kind': 2})
        options = analytics['filter_options']
        self.assertTrue(any(option['value'] == 2 for option in options['kind']))
        self.assertEqual(options['total'], 2)
        self.assertEqual(analytics['filters']['kind'], 2)

        rows = rec.filtered_rows(report, 'purchase', {'kind': 2})
        workbook = rec.build_analytics_workbook(analytics, rows)
        self.assertEqual(workbook.sheetnames, [
            'خلاصه', 'ماهانه', 'نوع معامله', 'نوع طرف معامله', 'وضعیت مودیان', 'طرف‌های برتر', 'ردیف‌ها',
        ])
        self.assertEqual(workbook['ردیف‌ها'].max_row, 2)
        self.assertEqual(workbook['نوع معامله'].cell(row=1, column=1).value, 'نوع')
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.assertGreater(buffer.tell(), 3000)

    def test_analytics_payload(self):
        report = self.make_report()
        sync_rows(report, [
            self.row(invoice_no='A1', amount=1000000, vat=100000, invoice_date='2026-04-10'),
            self.natural_row(invoice_no='A2', amount=500000, vat=50000, invoice_date='2026-05-10',
                             moadian_status='pending'),
            {'person_type': 'aggregate', 'amount': 250000, 'description': 'تجمیعی'},
        ])
        analytics = rec.build_analytics(report, 'purchase')
        self.assertEqual(analytics['totals']['count'], 3)
        self.assertEqual(analytics['totals']['amount'], 1750000.0)
        self.assertEqual(analytics['totals']['vat'], 150000.0)
        self.assertEqual(analytics['totals']['registered'], 1)
        self.assertEqual(analytics['quality']['aggregate_rows'], 1)
        self.assertTrue(analytics['by_month'])
        self.assertTrue(analytics['top_parties'])
        self.assertTrue(analytics['insights'])
        self.assertEqual(analytics['ruleset']['version'], rec.seasonal_rules.RULESET_VERSION)

